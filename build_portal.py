# -*- coding: utf-8 -*-
from __future__ import annotations

import html
import hashlib
import csv
import json
import math
import os
import re
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

import openpyxl
from PIL import Image, ImageOps
from build_guide import build_guide
from build_map_overlay import build_map_overlay


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT.parent / "игарка_реестр.xlsx"
OUTPUT = ROOT / "index.html"
PHOTOS_SOURCE = ROOT.parent / "Фото обнажений"
PHOTOS_FULL = ROOT / "photos"
PHOTOS_THUMBS = ROOT / "thumbs"
THIN_SECTIONS_SOURCE = ROOT.parent / "Шлифы"
THIN_SECTIONS_FULL = ROOT / "thin_sections"
THIN_SECTIONS_THUMBS = ROOT / "thin_section_thumbs"
PETROGRAPHY_SOURCE = ROOT.parent / "Claude outputs" / "shlify_igarka_opisaniya.md"
UPB_BUILDER = ROOT / "build_upb.py"
UPB_DATA = ROOT / "upb" / "portal_data.json"
GEOCHRONOLOGY_DATA = ROOT / "upb" / "geochronology_data.json"
MAJOR_CHEMISTRY = ROOT.parent / "Пасенко_Малышев_19_проб.xlsx"
TRACE_CHEMISTRY = ROOT.parent / "igarka_trace.xlsx"
PM_NORMALIZER = ROOT.parent / "PM_1989.csv"
REFERENCE_CHEMISTRY = Path(
    "/Users/sergeymalyshev/Library/CloudStorage/OneDrive-Personal/14-r/"
    "multel_dia_template/morb.csv"
)

# Literature localities from Kochnev et al. (2025). The article gives the
# 1812-2 GPS point and places 1812-1 40 m NW of it; the coordinate below is
# therefore a transparent geometric approximation, not a reported GPS fix.
LITERATURE_POINTS = [
    {
        "id": "KOCHNEV-1812-1",
        "sample": "1812-1",
        "author": "Кочнев и др., 2025",
        "formation": "силл в поле распространения губинской свиты",
        "material": "метадолеритовый силл",
        "lat": 67.180344,
        "lon": 86.579844,
        "coordinate_note": (
            "Расчётное положение: 40 м к СЗ от 1812-2 (67.18009° N, 86.58050° E), "
            "как указано в статье."
        ),
        "summary": "15 определений по 11 зернам; основной кластер около 835 млн лет.",
        "detail": (
            "Опубликованный concordia age 834,5 ± 4,2 млн лет (n=13). "
            "Авторы статьи трактуют цирконы как ксеногенные; это интерпретация, "
            "а не прямое измерение возраста внедрения тела."
        ),
        "dataset_id": "KOCHNEV-1812-1-SILL",
        "figure": "upb/kochnev2025_1812_1_concordia_wm.png",
        "audit": "upb/kochnev2025_1812_1_audit.csv",
        "doi": "https://doi.org/10.5800/GT-2025-16-1-0808",
    },
    {
        "id": "KOCHNEV-1812-2",
        "sample": "1812-2",
        "author": "Кочнев и др., 2022/2025",
        "formation": "губинская свита",
        "material": "розовато-серый косослоистый граувакково-кварцевый песчаник",
        "lat": 67.18009,
        "lon": 86.58050,
        "coordinate_note": "Точные координаты точки 1812-2, приведённые в статье 2025 года.",
        "summary": "130 анализов в приложении; 120 проходят опубликованный порог |D| ≤ 10%.",
        "detail": (
            "Опубликованный MDA 716 ± 10 млн лет по трём молодым зёрнам. "
            "Кнопка KDE открывает сопоставление 1812-2 с нашей пробой SM25-05."
        ),
        "dataset_id": "GUBINSKAYA-COMPARISON",
        "audit": "upb/kochnev2022_41812_audit.csv",
        "doi": "https://doi.org/10.1016/j.precamres.2022.106877",
    },
    {
        "id": "KOCHNEV-2007-15",
        "sample": "2007-15",
        "author": "Кочнев и др., 2025",
        "formation": "лудовская свита",
        "material": "крупнозернистый полимиктовый песчаник нижней части свиты",
        "lat": 67.80394,
        "lon": 86.42067,
        "coordinate_note": (
            "Координаты основания типового разреза из статьи; образец отобран "
            "в 83 м выше подошвы обнажённого разреза."
        ),
        "summary": (
            "121 определение; самое молодое 1551 ± 27 млн лет, остальные — "
            "палеопротерозойские и архейские."
        ),
        "detail": "Основные опубликованные пики: около 2,76, 2,57, 2,05, 1,93 и 1,87 млрд лет.",
        "dataset_id": "KOCHNEV-2007-15-LUDY",
        "audit": "upb/kochnev2025_2007_15_audit.csv",
        "doi": "https://doi.org/10.5800/GT-2025-16-1-0808",
    },
]


def dms_to_decimal(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).replace("\u00a0", " ").strip()
    match = re.search(
        r"(\d+(?:[.,]\d+)?)\s*°\s*(\d+(?:[.,]\d+)?)\s*['′]\s*(\d+(?:[.,]\d+)?)\s*[\"″]?\s*([NSEWСЮВЗ]?)",
        text,
        re.IGNORECASE,
    )
    if not match:
        return None
    degrees, minutes, seconds = (float(v.replace(",", ".")) for v in match.groups()[:3])
    result = degrees + minutes / 60 + seconds / 3600
    if match.group(4).upper() in {"S", "W", "Ю", "З"}:
        result = -result
    return result


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


TASK_LABELS = {
    "thin_section": "Шлиф",
    "chemistry": "Химсостав",
    "upb": "U–Pb",
    "sm_nd": "Sm–Nd",
    "aft": "Треки AFT",
}


def task_state(value: str, planned: bool) -> str:
    """Translate free-form registry wording into a transparent workflow stage."""
    if not planned:
        return "not_planned"
    value = value.lower()
    if any(word in value for word in ("сделано", "отработано", "готово")) or re.search(
        r"(^|[\s,;])есть([\s,;]|$)", value
    ):
        return "done"
    if any(word in value for word in ("передано", "отдано", "в лаб", "анализ d", "анализ ap")):
        return "in_lab"
    if any(word in value for word in ("выделено", "навеск", "подготовлен")):
        return "prepared"
    if any(word in value for word in ("выделение", "дробление", "сепарац")):
        return "preparing"
    return "planned"


def workflow(record: dict[str, object]) -> tuple[list[dict[str, str]], str]:
    general = " · ".join(
        value for value in (
            text(record.get("status_raw")),
            text(record.get("description")),
            text(record.get("comment")),
        ) if value
    )
    low = general.lower()
    rules = {
        "thin_section": ("шлиф",),
        "chemistry": ("хим",),
        "upb": (
            "u-pb", "u–pb", "геохрон", "циркон", "dzr", "d-zr", "бадделеит", "zr ",
            "апатит", "анализ ap", " ap ",
        ),
        "sm_nd": ("sm-nd", "sm–nd", "sm nd"),
        # Apatite separation alone does not imply fission-track analysis:
        # apatite is also prepared for U–Pb dating. AFT is planned only when
        # the registry's dedicated AFT cell is filled or tracks/AFT are named
        # explicitly in the status, description, or comment.
        "aft": ("aft", "трек"),
    }
    tasks: list[dict[str, str]] = []
    for key, label in TASK_LABELS.items():
        source_key = "upb_status" if key == "upb" else key
        direct = text(record.get(source_key))
        mentioned = any(token in f" {low} " for token in rules[key])
        planned = bool(direct) or mentioned
        relevant = direct
        if mentioned:
            relevant = " · ".join(part for part in (direct, text(record.get("status_raw")), text(record.get("comment"))) if part)
        tasks.append({
            "key": key,
            "label": label,
            "state": task_state(relevant, planned),
            "detail": relevant or ("Запланировано по описанию образца" if mentioned else "Не запланировано"),
        })

    planned_tasks = [task for task in tasks if task["state"] != "not_planned"]
    if not planned_tasks:
        status = "no_plan"
    elif all(task["state"] == "done" for task in planned_tasks):
        status = "completed"
    elif any(task["state"] == "in_lab" for task in planned_tasks):
        status = "lab"
    elif any(task["state"] in {"preparing", "prepared"} for task in planned_tasks):
        status = "preparation"
    else:
        status = "planned"
    return tasks, status


def photo_sort_key(path: Path) -> tuple[str, int, str]:
    stem = path.stem
    point, separator, suffix = stem.partition("_")
    number = int(suffix) if separator and suffix.isdigit() else 0
    return point, number, path.name.lower()


def safe_photo_name(path: Path) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", path.stem.replace("*", "_star")) + ".jpg"


def safe_thin_section_name(path: Path) -> str:
    normalized = unicodedata.normalize("NFC", path.name)
    prefix = re.sub(r"[^A-Za-z0-9._-]+", "_", normalized.split("-анализ", 1)[0])[:48].strip("_")
    digest = hashlib.sha1(normalized.encode("utf-8")).hexdigest()[:10]
    return f"{prefix or 'thin_section'}_{digest}.jpg"


def open_photo(path: Path) -> Image.Image:
    try:
        return Image.open(path)
    except Exception:
        if path.suffix.lower() not in {".heic", ".heif"}:
            raise
        with tempfile.TemporaryDirectory(prefix="igarka_photo_") as temporary:
            converted = Path(temporary) / "converted.jpg"
            subprocess.run(
                ["/usr/local/bin/magick", str(path), "-auto-orient", str(converted)],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            image = Image.open(converted)
            image.load()
            return image


def prepare_photo(
    source: Path,
    full_path: Path,
    thumb_path: Path,
    full_size: int = 1800,
    full_quality: int = 88,
) -> None:
    if (
        full_path.exists()
        and thumb_path.exists()
        and full_path.stat().st_mtime >= source.stat().st_mtime
        and thumb_path.stat().st_mtime >= source.stat().st_mtime
    ):
        return
    image = ImageOps.exif_transpose(open_photo(source)).convert("RGB")
    full = image.copy()
    full.thumbnail((full_size, full_size), Image.LANCZOS)
    full.save(full_path, "JPEG", quality=full_quality, optimize=True, progressive=True)
    thumb = image.copy()
    thumb.thumbnail((420, 420), Image.LANCZOS)
    thumb.save(thumb_path, "JPEG", quality=82, optimize=True, progressive=True)


def attach_photos(records: list[dict[str, object]]) -> tuple[int, list[str]]:
    PHOTOS_FULL.mkdir(parents=True, exist_ok=True)
    PHOTOS_THUMBS.mkdir(parents=True, exist_ok=True)
    known_points = {str(record["sample"]).split("/")[0] for record in records}
    photos_by_point: dict[str, list[dict[str, str]]] = {}
    unmatched: list[str] = []
    source_files = sorted(
        [
            path
            for path in PHOTOS_SOURCE.iterdir()
            if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".heic", ".heif"}
        ],
        key=photo_sort_key,
    )
    for source in source_files:
        point = source.stem.split("_", 1)[0]
        if point not in known_points:
            unmatched.append(source.name)
            continue
        output_name = safe_photo_name(source)
        full_path = PHOTOS_FULL / output_name
        thumb_path = PHOTOS_THUMBS / output_name
        prepare_photo(source, full_path, thumb_path)
        photos_by_point.setdefault(point, []).append(
            {
                "thumb": f"thumbs/{output_name}",
                "full": f"photos/{output_name}",
                "name": source.name,
            }
        )
    for record in records:
        point = str(record["sample"]).split("/")[0]
        record["photo_point"] = point
        record["photos"] = photos_by_point.get(point, [])
    return sum(len(items) for items in photos_by_point.values()), unmatched


def thin_section_mode(filename: str) -> str:
    lowered = filename.casefold()
    if "николи х" in lowered or "скрещ" in lowered or "xpl" in lowered:
        return "XPL · скрещённые николи"
    if "николи пар" in lowered or "паралл" in lowered or "ppl" in lowered:
        return "PPL · параллельные николи"
    if "безанализ" in lowered or "дбезанализ" in lowered:
        return "Без анализатора"
    if "анализатор" in lowered:
        return "С анализатором"
    return "Микрофото"


def attach_thin_sections(records: list[dict[str, object]]) -> tuple[int, list[str]]:
    THIN_SECTIONS_FULL.mkdir(parents=True, exist_ok=True)
    THIN_SECTIONS_THUMBS.mkdir(parents=True, exist_ok=True)
    samples = [str(record["sample"]) for record in records]
    aliases: list[tuple[str, str]] = []
    for sample in samples:
        aliases.append((unicodedata.normalize("NFC", sample.upper().replace("/", "-")), sample))
        if "*" in sample:
            aliases.append(
                (unicodedata.normalize("NFC", sample.upper().replace("*", " СО ЗВЕЗДОЧКОЙ")), sample)
            )
    aliases.sort(key=lambda item: len(item[0]), reverse=True)
    thin_sections_by_sample: dict[str, list[dict[str, str]]] = {}
    unmatched: list[str] = []
    expected_outputs: set[str] = set()
    sources = sorted(
        [
            path
            for path in THIN_SECTIONS_SOURCE.iterdir()
            if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}
        ],
        key=lambda path: path.name.casefold(),
    )
    for source in sources:
        normalized = unicodedata.normalize("NFC", source.stem.upper())
        matched_sample = None
        for alias, sample in aliases:
            if normalized == alias or (
                normalized.startswith(alias) and normalized[len(alias) : len(alias) + 1] in {"-", "_"}
            ):
                matched_sample = sample
                break
        if matched_sample is None:
            unmatched.append(source.name)
            continue
        output_name = safe_thin_section_name(source)
        expected_outputs.add(output_name)
        full_path = THIN_SECTIONS_FULL / output_name
        thumb_path = THIN_SECTIONS_THUMBS / output_name
        prepare_photo(source, full_path, thumb_path, full_size=2560, full_quality=90)
        thin_sections_by_sample.setdefault(matched_sample, []).append(
            {
                "thumb": f"thin_section_thumbs/{output_name}",
                "full": f"thin_sections/{output_name}",
                "name": source.name,
                "mode": thin_section_mode(source.name),
            }
        )
    for record in records:
        record["thin_section_photos"] = thin_sections_by_sample.get(str(record["sample"]), [])
    for output_directory in (THIN_SECTIONS_FULL, THIN_SECTIONS_THUMBS):
        for generated in output_directory.glob("*.jpg"):
            if generated.name not in expected_outputs:
                generated.unlink()
    return sum(len(items) for items in thin_sections_by_sample.values()), unmatched


def clean_markdown_text(value: str) -> str:
    cleaned = value.replace("**", "").replace("*", "").strip()
    cleaned = re.sub(r"\[([^\]]+)\]", r"\1", cleaned)
    return re.sub(r"[ \t]+", " ", cleaned)


def attach_petrography(records: list[dict[str, object]]) -> tuple[int, list[str]]:
    source = PETROGRAPHY_SOURCE.read_text(encoding="utf-8")
    heading_pattern = re.compile(
        r"^##\s+(SM\d{2}-\d{2}(?:/\d+|\*)?)\s+—\s+(.+?)\s*$", re.MULTILINE
    )
    matches = list(heading_pattern.finditer(source))
    descriptions_by_sample: dict[str, list[dict[str, str]]] = {}
    unmatched: list[str] = []
    known_samples = {str(record["sample"]) for record in records}

    for index, match in enumerate(matches):
        section_id = match.group(1)
        target_sample = "SM25-08" if section_id.startswith("SM25-08/") else section_id
        if target_sample not in known_samples:
            unmatched.append(section_id)
            continue
        # Stop at the next level-two heading, including editorial sections such
        # as “Образцы без шлифа”; otherwise their text leaks into the final
        # sample description when that sample is last in the file.
        following_heading = re.search(r"^##\s+", source[match.end() :], re.MULTILINE)
        end = (
            match.end() + following_heading.start()
            if following_heading is not None
            else len(source)
        )
        raw_section = source[match.end() : end].strip()
        paragraphs = [paragraph.strip() for paragraph in re.split(r"\n\s*\n", raw_section) if paragraph.strip()]
        body: list[str] = []
        notes: list[str] = []
        for paragraph in paragraphs:
            if paragraph == "---":
                continue
            if paragraph.startswith("**Локация"):
                continue
            if paragraph.startswith("*Примечание") or paragraph.startswith("**Комментарий"):
                notes.append(clean_markdown_text(paragraph))
                continue
            if paragraph.startswith("*(У аспирантки"):
                notes.append(clean_markdown_text(paragraph))
                continue
            body.append(clean_markdown_text(paragraph))
        descriptions_by_sample.setdefault(target_sample, []).append(
            {
                "section": section_id,
                "title": clean_markdown_text(match.group(2)),
                "description": "\n\n".join(body),
                "note": "\n\n".join(notes),
            }
        )

    for record in records:
        record["petrography"] = descriptions_by_sample.get(str(record["sample"]), [])
    return sum(len(items) for items in descriptions_by_sample.values()), unmatched


def attach_upb(records: list[dict[str, object]]) -> int:
    if not UPB_DATA.exists():
        return 0
    results = json.loads(UPB_DATA.read_text(encoding="utf-8"))
    literature_cards = {}
    if GEOCHRONOLOGY_DATA.exists():
        literature_cards = json.loads(GEOCHRONOLOGY_DATA.read_text(encoding="utf-8")).get(
            "literature_cards", {}
        )
    attached = 0
    for record in records:
        sample = str(record["sample"])
        record["upb"] = results.get(sample)
        record["literature_geochronology"] = literature_cards.get(sample)
        if record["upb"]:
            attached += 1
    return attached


def chemistry_sample_id(value: object) -> str:
    return re.sub(r"\s+", "", text(value)).upper()


def read_geochemical_csv(path: Path) -> tuple[list[str], list[dict[str, object]]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    elements = [text(value) for value in rows[0][1:]]
    datasets: list[dict[str, object]] = []
    for row in rows[1:]:
        if not text(row[0]):
            continue
        datasets.append(
            {
                "name": text(row[0]),
                "values": {
                    element: float(value)
                    for element, value in zip(elements, row[1:])
                    if element and text(value)
                },
            }
        )
    return elements, datasets


def geochemical_context(record: dict[str, object]) -> str:
    combined = " ".join(
        text(record.get(key)).casefold()
        for key in ("location", "description")
    )
    if "силл" in combined:
        return "sill"
    if "дайк" in combined:
        return "igarka_dyke"
    if "туф" in combined:
        return "tuff"
    return "other"


def attach_geochemistry(records: list[dict[str, object]]) -> int:
    major_book = openpyxl.load_workbook(MAJOR_CHEMISTRY, data_only=True, read_only=True)
    major_sheet = major_book["Лист1"]
    major_headers = [text(cell.value).replace(" (%)", "") for cell in major_sheet[4]]
    majors: dict[str, dict[str, float]] = {}
    for row in major_sheet.iter_rows(min_row=5, values_only=True):
        sample_id = chemistry_sample_id(row[0])
        if not sample_id:
            continue
        majors[sample_id] = {
            major_headers[index]: float(value)
            for index, value in enumerate(row[1:], start=1)
            if value is not None
        }

    trace_book = openpyxl.load_workbook(TRACE_CHEMISTRY, data_only=True, read_only=True)
    trace_sheet = trace_book["Sheet1"]
    trace_headers = [text(cell.value) for cell in trace_sheet[1]]
    traces: dict[str, dict[str, object]] = {}
    for row in trace_sheet.iter_rows(min_row=2, values_only=True):
        sample_id = chemistry_sample_id(row[0])
        if not sample_id or not text(row[1]):
            continue
        traces[sample_id] = {
            "group": text(row[1]),
            "values": {
                trace_headers[index]: float(value)
                for index, value in enumerate(row[2:], start=2)
                if trace_headers[index] and value is not None
            },
        }

    attached = 0
    for record in records:
        sample_id = chemistry_sample_id(record["sample"])
        major = majors.get(sample_id)
        trace = traces.get(sample_id)
        record["geochem"] = (
            {
                "major": major or {},
                "trace": trace["values"] if trace else {},
                "group": trace["group"] if trace else "",
                "context": geochemical_context(record),
                "approximate": ["V", "Cr", "Co", "Ni", "Cu", "Zn", "Mo", "Sn", "Tl"],
            }
            if major or trace
            else None
        )
        if record["geochem"]:
            attached += 1
    return attached


def load_records() -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    sheet = workbook["Sheet2"]
    source_rows: list[tuple[int, int | None, tuple[object, ...]]] = []
    current_year = None
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if text(row[4]) == "ИГАРКА 2025":
            current_year = 2025
            continue
        if text(row[4]) == "ИГАРКА 2026":
            current_year = 2026
            continue
        if text(row[1]).startswith("SM"):
            source_rows.append((row_number, current_year, row))

    sample_ids = {text(row[1]) for _, _, row in source_rows}
    point_header_ids = {
        sample_id.split("/", 1)[0]
        for sample_id in sample_ids
        if "/" in sample_id and sample_id.split("/", 1)[0] in sample_ids
    }
    point_headers = {
        text(row[1]): {"row": row_number, "year": row_year, "values": row}
        for row_number, row_year, row in source_rows
        if text(row[1]) in point_header_ids
    }

    records: list[dict[str, object]] = []
    parent_coordinates: dict[str, tuple[float, float]] = {}

    for row_number, year, row in source_rows:
        sample = text(row[1])
        if sample in point_header_ids:
            header_lat = dms_to_decimal(row[2])
            header_lon = dms_to_decimal(row[3])
            if header_lat is not None and header_lon is not None:
                parent_coordinates[sample] = (header_lat, header_lon)
            continue

        lat = dms_to_decimal(row[2])
        lon = dms_to_decimal(row[3])
        base_sample = sample.split("/")[0]
        point_header = point_headers.get(base_sample)
        point_values = point_header["values"] if point_header else None
        inherited = False
        if lat is not None and lon is not None:
            parent_coordinates[base_sample] = (lat, lon)
        elif base_sample in parent_coordinates:
            lat, lon = parent_coordinates[base_sample]
            inherited = True

        record: dict[str, object] = {
            "order": len(records) + 1,
            "row": row_number,
            "year": year,
            "sample": sample,
            "point_id": base_sample,
            "point_row": point_header["row"] if point_header else row_number,
            "point_description": text(point_values[4]) if point_values else "",
            "point_note": text(point_values[5]) if point_values else "",
            "lat": lat,
            "lon": lon,
            "coords_inherited": inherited,
            "location": text(row[4]),
            "description": text(row[5]),
            "status_raw": text(row[6]),
            "thin_section": text(row[7]),
            "chemistry": text(row[8]),
            "aft": text(row[9]),
            "sm_nd": text(row[10]),
            "upb_status": text(row[11]),
            "comment": text(row[12]),
        }
        record["workflow_tasks"], record["workflow_status"] = workflow(record)
        records.append(record)

    return records


def build_html(
    records: list[dict[str, object]],
    spider_elements: list[str],
    primitive_mantle: dict[str, object],
    references: list[dict[str, object]],
    geochronology: dict[str, object],
) -> str:
    payload = json.dumps(records, ensure_ascii=False, separators=(",", ":"))
    spider_payload = json.dumps(spider_elements, ensure_ascii=False, separators=(",", ":"))
    primitive_payload = json.dumps(primitive_mantle, ensure_ascii=False, separators=(",", ":"))
    references_payload = json.dumps(references, ensure_ascii=False, separators=(",", ":"))
    geochronology_payload = json.dumps(geochronology, ensure_ascii=False, separators=(",", ":"))
    literature_points_payload = json.dumps(
        LITERATURE_POINTS, ensure_ascii=False, separators=(",", ":")
    )
    template = r'''<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<script>if(new URLSearchParams(location.search).has('mapshot'))document.documentElement.classList.add('mapshot');</script>
<title>Игарка — реестр образцов</title>
<style>
:root {
  --panel:#fbfcfa; --paper:#ffffff; --line:#d7ddd8; --text:#17211b; --muted:#68736c;
  --done:#23865b; --progress:#3278bd; --prepared:#d08a20; --empty:#8b948f; --accent:#315f46;
}
* { box-sizing:border-box; }
html,body { margin:0; width:100%; height:100%; overflow:hidden; }
body { font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif; color:var(--text); background:#e8ece9; }
button,input,select { font:inherit; }
.app { display:grid; grid-template-columns:300px minmax(420px,1fr) 350px; height:100vh; min-height:560px; overflow:hidden; }
html.mapshot .app { grid-template-columns:1fr; }
html.mapshot .left,html.mapshot .right,html.mapshot .map-toolbar,html.mapshot .map-note { display:none; }
html.mapshot .map-wrap { height:100vh; }
html.mapshot #kochnevLegend { display:none; }
.panel { min-width:0; min-height:0; height:100vh; overflow:hidden; background:var(--panel); display:flex; flex-direction:column; z-index:3; }
.left { border-right:1px solid var(--line); }
.right { border-left:1px solid var(--line); }
.panel-head { padding:16px 16px 12px; border-bottom:1px solid var(--line); background:rgba(255,255,255,.92); }
.eyebrow { color:var(--accent); font-size:10px; font-weight:800; letter-spacing:.08em; text-transform:uppercase; }
h1 { margin:5px 0 4px; font-size:21px; line-height:1.12; }
.hint { color:var(--muted); font-size:11.5px; line-height:1.38; }
.summary { display:grid; grid-template-columns:repeat(5,1fr); gap:4px; margin-top:12px; }
.metric { min-width:0; padding:7px 4px; border:1px solid var(--line); background:#fff; border-radius:7px; text-align:center; }
.metric strong { display:block; font-size:16px; line-height:1; }
.metric span { display:block; margin-top:4px; color:var(--muted); font-size:8.5px; text-transform:uppercase; }
.controls { display:grid; gap:7px; padding:10px; border-bottom:1px solid var(--line); }
.controls-row { display:grid; grid-template-columns:1fr 1fr; gap:7px; }
input,select { width:100%; border:1px solid #cbd3cd; border-radius:7px; padding:8px 9px; color:var(--text); background:white; font-size:12px; outline:none; }
input:focus,select:focus { border-color:#6f9a80; box-shadow:0 0 0 2px rgba(49,95,70,.12); }
.list-summary { padding:8px 11px 4px; color:var(--muted); font-size:10px; }
.list { min-height:0; flex:1; overflow:auto; padding:5px 7px 12px; }
.item { margin-bottom:5px; padding:8px 9px; border:1px solid transparent; border-radius:8px; background:#fff; cursor:pointer; box-shadow:0 1px 0 rgba(24,43,31,.05); }
.item:hover { border-color:#adc0b3; }
.item.active { border-color:#4f8968; box-shadow:0 0 0 1px #4f8968; }
.item-top { display:flex; justify-content:space-between; gap:8px; align-items:center; }
.sample { font-weight:800; font-size:12.5px; }
.year { color:var(--muted); font-size:9.5px; }
.item-location { margin-top:4px; color:#465149; font-size:11px; line-height:1.25; }
.status-line { display:flex; align-items:center; gap:6px; margin-top:6px; }
.status-name { color:var(--muted); font-size:9.5px; flex:1; }
.task-chips { display:flex; flex-wrap:wrap; gap:3px; margin-top:6px; }
.task-chip,.task-state { display:inline-flex; align-items:center; border:1px solid #ccd4cf; border-radius:99px; padding:2px 5px; color:#59645d; background:#f4f6f4; font-size:8.5px; line-height:1.2; }
.task-chip.done,.task-state.done { border-color:#9bc7ae; color:#1f704d; background:#e9f5ee; }
.task-chip.in_lab,.task-state.in_lab { border-color:#99b9db; color:#245f99; background:#eaf2fa; }
.task-chip.preparing,.task-chip.prepared,.task-state.preparing,.task-state.prepared { border-color:#dfbf83; color:#8b5b0f; background:#fbf3e4; }
.task-chip.planned,.task-state.planned { border-color:#c8cdd0; color:#646d72; background:#f3f4f5; }
.task-state.not_planned { color:#929994; background:transparent; }
.map-wrap { position:relative; min-width:0; min-height:0; height:100vh; overflow:hidden; background:#1b2626; }
#map { width:100%; height:100%; display:block; cursor:grab; }
#map.dragging { cursor:grabbing; }
.map-toolbar { position:absolute; top:12px; left:12px; display:flex; flex-wrap:wrap; max-width:calc(100% - 24px); gap:5px; padding:5px; border:1px solid rgba(255,255,255,.45); border-radius:9px; background:rgba(25,33,29,.76); backdrop-filter:blur(8px); box-shadow:0 8px 24px rgba(0,0,0,.18); }
.map-toolbar button { min-width:31px; height:30px; padding:0 9px; border:1px solid rgba(255,255,255,.22); border-radius:6px; color:white; background:rgba(255,255,255,.12); cursor:pointer; }
.map-toolbar button:hover { background:rgba(255,255,255,.24); }
.map-toolbar button[aria-pressed="true"] { border-color:rgba(255,255,255,.75); background:rgba(255,255,255,.28); }
.map-note { position:absolute; top:52px; right:12px; padding:7px 9px; border-radius:7px; color:rgba(255,255,255,.9); background:rgba(25,33,29,.72); font-size:9.5px; }
.map-layer-control { position:absolute; left:12px; top:112px; width:235px; max-width:calc(100% - 24px); padding:11px 12px; border:1px solid rgba(255,255,255,.4); border-radius:9px; color:white; background:rgba(25,33,29,.85); backdrop-filter:blur(8px); font-size:12px; }
.map-layer-control label { display:flex; align-items:center; justify-content:space-between; gap:8px; }
.map-layer-control .map-layer-heading { justify-content:flex-start; font-weight:700; }
.map-layer-control input[type="checkbox"] { width:16px; height:16px; margin:0; padding:0; accent-color:#b9d4bc; }
.map-layer-control input[type="range"] { display:block; width:100%; height:23px; padding:0; margin:6px 0 0; accent-color:#c1ddc5; background:transparent; border:0; box-shadow:none; cursor:pointer; }
.map-layer-slider { margin-top:12px; }
.map-layer-slider.inactive { opacity:.45; }
.map-layer-status { margin-top:6px; font-size:10px; color:#e1e9df; }
.map-layer-status:empty { display:none; }
html.mapshot .map-layer-control { display:none; }
.map-legend { position:absolute; top:84px; right:12px; display:grid; gap:6px; padding:8px 10px; border-radius:7px; color:rgba(255,255,255,.94); background:rgba(25,33,29,.72); font-size:9.5px; }
.map-legend-item { display:flex; align-items:center; gap:7px; }
.legend-pin { position:relative; width:12px; height:12px; border:1.5px solid white; border-radius:50% 50% 50% 0; transform:rotate(-45deg); box-shadow:0 2px 4px rgba(0,0,0,.35); }
.legend-pin::after { content:""; position:absolute; left:3px; top:3px; width:3px; height:3px; border-radius:50%; background:white; }
.legend-pin.y2025 { background:#19765f; }
.legend-pin.y2026 { background:#c47725; }
.legend-pin.kochnev { background:#6b5a92; }
.popup { position:absolute; display:none; min-width:190px; max-width:270px; padding:8px; border:1px solid #bdc7c0; border-radius:9px; background:white; box-shadow:0 12px 34px rgba(0,0,0,.24); z-index:7; }
.popup.visible { display:block; }
.popup-title { margin-bottom:6px; color:var(--muted); font-size:10px; }
.popup button { display:block; width:100%; margin-top:3px; padding:6px 7px; border:1px solid var(--line); border-radius:6px; background:white; text-align:left; cursor:pointer; font-size:11px; }
.popup button:hover,.popup button.active { border-color:#5c9270; background:#edf5f0; }
.detail-scroll { min-height:0; flex:1; overflow:auto; padding:13px 14px 24px; }
.hero-status { display:flex; align-items:flex-start; gap:9px; margin-bottom:11px; padding:10px; border:1px solid var(--line); border-radius:9px; background:white; }
.hero-status strong { display:block; font-size:12px; }
.hero-status span { color:var(--muted); font-size:10px; }
.section { margin-top:15px; }
.section-title { margin-bottom:7px; color:var(--muted); font-size:9.5px; font-weight:800; letter-spacing:.06em; text-transform:uppercase; }
.field { margin-top:8px; font-size:12px; line-height:1.42; white-space:pre-wrap; }
.field b { display:block; margin-bottom:2px; color:var(--muted); font-size:9.5px; font-weight:700; text-transform:uppercase; }
.tracks { display:grid; gap:6px; }
.track { display:grid; grid-template-columns:92px 72px 1fr; gap:7px; align-items:start; padding:7px 8px; border:1px solid var(--line); border-radius:7px; background:white; font-size:11px; }
.track-name { color:var(--muted); }
.coord { font-variant-numeric:tabular-nums; }
.source-note { margin-top:15px; padding-top:10px; border-top:1px solid var(--line); color:var(--muted); font-size:9.5px; line-height:1.4; }
.photo-grid { display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:6px; }
.photo-thumb-button { position:relative; display:block; width:100%; aspect-ratio:1/1; padding:0; overflow:hidden; border:1px solid var(--line); border-radius:7px; background:#e8ece9; cursor:zoom-in; }
.photo-thumb-button:hover { border-color:#5c9270; box-shadow:0 0 0 1px #5c9270; }
.photo-thumb-button img { width:100%; height:100%; display:block; object-fit:cover; }
.photo-number { position:absolute; right:4px; bottom:4px; min-width:18px; padding:2px 4px; border-radius:99px; color:white; background:rgba(18,27,22,.72); font-size:9px; text-align:center; }
.photo-mode { position:absolute; left:4px; bottom:4px; max-width:calc(100% - 30px); padding:2px 5px; overflow:hidden; border-radius:99px; color:white; background:rgba(18,27,22,.72); font-size:8px; line-height:1.2; text-overflow:ellipsis; white-space:nowrap; }
.photo-empty { padding:10px; border:1px dashed #cbd3cd; border-radius:7px; color:var(--muted); font-size:10.5px; text-align:center; }
.petro-list { display:grid; gap:8px; }
.petro-card { padding:10px; border:1px solid var(--line); border-radius:8px; background:white; }
.petro-heading { margin-bottom:7px; font-size:11.5px; font-weight:800; line-height:1.3; }
.petro-heading span { display:block; margin-bottom:2px; color:var(--accent); font-size:9px; letter-spacing:.04em; }
.petro-text { font-size:11px; line-height:1.48; white-space:pre-wrap; }
.petro-note { margin-top:8px; padding:7px 8px; border-left:3px solid var(--prepared); border-radius:0 6px 6px 0; color:#59645d; background:#faf5ea; font-size:10px; line-height:1.42; white-space:pre-wrap; }
.petro-source { margin-top:7px; color:var(--muted); font-size:8.5px; }
.upb-card { overflow:hidden; border:1px solid var(--line); border-radius:9px; background:white; }
.upb-figure { display:block; width:100%; padding:0; border:0; background:#f7f9f7; cursor:zoom-in; }
.upb-figure img { display:block; width:100%; height:auto; }
.upb-summary { padding:9px 10px 10px; font-size:10.5px; line-height:1.45; }
.upb-interpretation { margin-top:7px; padding-top:7px; border-top:1px solid var(--line); color:#4f5b55; }
.upb-links { display:flex; gap:8px; margin-top:6px; }
.upb-links a { color:var(--accent); font-size:9.5px; text-decoration:none; }
.upb-links a:hover { text-decoration:underline; }
.literature-card { padding:10px; border:1px solid #cad6cf; border-left:3px solid #6b5a92; border-radius:8px; background:#faf9fc; font-size:10.5px; line-height:1.45; }
.literature-card strong { display:block; margin-bottom:5px; font-size:11.5px; }
.literature-card p { margin:5px 0; }
.literature-source { margin-top:7px; color:var(--muted); font-size:9px; }
.literature-actions { display:flex; flex-wrap:wrap; gap:7px; margin-top:8px; }
.literature-actions button,.literature-actions a { padding:6px 8px; border:1px solid #c8c0d8; border-radius:6px; color:#564775; background:white; font-size:9.5px; text-decoration:none; cursor:pointer; }
.geo-modal { position:fixed; inset:0; z-index:46; display:none; grid-template-rows:auto minmax(0,1fr); background:#eef1ef; }
.geo-modal.visible { display:grid; }
.geo-toolbar { display:flex; flex-wrap:wrap; align-items:center; gap:12px 18px; padding:11px 16px; border-bottom:1px solid var(--line); background:white; }
.geo-heading { min-width:230px; margin-right:auto; }
.geo-heading strong { display:block; font-size:15px; }
.geo-heading span { color:var(--muted); font-size:9.5px; }
.geo-control { display:grid; grid-template-columns:auto 180px; align-items:center; gap:4px 9px; color:var(--muted); font-size:9.5px; }
.geo-control span { color:var(--text); font-weight:800; }
.geo-control input { grid-column:1/-1; width:230px; padding:0; }
.geo-close { padding:8px 11px; border:1px solid #cbd3cd; border-radius:6px; background:white; cursor:pointer; }
.geo-scroll { min-height:0; overflow:auto; padding:14px; }
.geo-note { width:min(100%,1440px); margin:0 auto 10px; color:#4f5b55; font-size:10px; line-height:1.45; }
.geo-grid { display:grid; grid-template-columns:repeat(2,minmax(0,700px)); justify-content:center; align-items:start; gap:14px; }
.geo-card { min-width:0; padding:11px; border:1px solid var(--line); border-radius:9px; background:white; }
.geo-card.focused { border-color:#6f9a80; box-shadow:0 0 0 1px #6f9a80; }
.geo-card-head { display:flex; justify-content:space-between; gap:12px; margin-bottom:4px; }
.geo-card-head strong { display:block; font-size:12px; }
.geo-card-head span { color:var(--muted); font-size:9px; text-align:right; }
.geo-chart { min-height:280px; }
.geo-young-chart { min-height:235px; }
.geo-chart svg { display:block; width:100%; height:auto; }
.geo-young-chart svg { display:block; width:100%; height:auto; }
.geo-chart-label { display:flex; justify-content:space-between; gap:10px; margin:8px 0 1px; color:#465149; font-size:9.5px; font-weight:800; }
.geo-chart-label span { color:var(--muted); font-weight:500; }
.geo-series-legend { display:flex; flex-wrap:wrap; gap:6px 12px; margin:2px 0 5px; color:#4f5b55; font-size:9px; }
.geo-series-legend span { display:inline-flex; align-items:center; gap:5px; }
.geo-series-legend label { display:inline-flex; align-items:center; gap:5px; padding:4px 7px; border:1px solid #cbd3cd; border-radius:6px; cursor:pointer; }
.geo-series-legend input { margin:0; }
.geo-series-legend i { width:18px; height:3px; border-radius:3px; }
.geo-peaks { display:flex; flex-wrap:wrap; gap:5px; margin-top:5px; }
.geo-peak { padding:5px 7px; border:1px solid #cbd3cd; border-radius:99px; background:white; color:#465149; font-size:9px; cursor:pointer; }
.geo-peak.active { color:white; border-color:var(--accent); background:var(--accent); }
.geo-stats { min-height:42px; margin-top:7px; padding-top:7px; border-top:1px solid var(--line); color:#465149; font-size:9.5px; line-height:1.45; }
.geo-source { margin-top:5px; color:var(--muted); font-size:8.5px; }
.geo-young-open { margin-top:7px; padding:7px 9px; border:1px solid #b8c7bf; border-radius:6px; color:#315f50; background:#f5faf7; font-size:9.5px; cursor:pointer; }
.young-modal { --ym-bg:#151718; --ym-panel:#1c1f21; --ym-text:#f2f2f2; --ym-muted:#999d9f; --ym-grid:#34383b; position:fixed; inset:0; z-index:49; display:none; grid-template-rows:auto minmax(0,1fr); color:var(--ym-text); background:var(--ym-bg); }
.young-modal.visible { display:grid; }
.young-modal.light { --ym-bg:#eef1ef; --ym-panel:#fff; --ym-text:#26322d; --ym-muted:#68736c; --ym-grid:#dfe4e0; }
.young-toolbar { display:flex; flex-wrap:wrap; align-items:center; gap:8px 12px; padding:11px 16px; border-bottom:1px solid var(--ym-grid); background:var(--ym-panel); }
.young-toolbar strong { margin-right:auto; font-size:14px; }
.young-series-buttons { display:flex; gap:5px; }
.young-toolbar button { padding:7px 9px; border:1px solid var(--ym-grid); border-radius:6px; color:var(--ym-text); background:transparent; cursor:pointer; }
.young-toolbar button.active { color:#101716; border-color:#76b9ee; background:#76b9ee; }
.young-content { min-height:0; overflow:auto; padding:18px; }
.young-panel { width:min(100%,1280px); margin:auto; padding:14px 16px 10px; background:var(--ym-panel); }
.young-meta { display:flex; flex-wrap:wrap; gap:8px 24px; margin-bottom:9px; color:var(--ym-muted); font-size:11px; }
.young-chart svg { display:block; width:100%; height:auto; }
.young-note { margin-top:8px; color:var(--ym-muted); font-size:9.5px; line-height:1.45; }
.upb-modal { position:fixed; inset:0; z-index:47; display:none; grid-template-rows:auto minmax(0,1fr); background:#eef1ef; }
.upb-modal.visible { display:grid; }
.upb-toolbar { display:flex; flex-wrap:wrap; align-items:end; gap:10px 16px; padding:11px 16px; border-bottom:1px solid var(--line); background:white; }
.upb-heading { min-width:180px; margin-right:auto; }
.upb-heading strong { display:block; font-size:14px; }
.upb-heading span { color:var(--muted); font-size:9.5px; }
.upb-control { display:grid; gap:3px; color:var(--muted); font-size:9px; }
.upb-control select { min-width:210px; padding:6px 8px; border:1px solid #cbd3cd; border-radius:6px; background:white; }
.upb-slider-row { display:flex; align-items:center; gap:7px; }
.upb-slider-row input { width:160px; }
.upb-preset,.upb-close,.upb-zoom button { padding:7px 9px; border:1px solid #cbd3cd; border-radius:6px; background:white; cursor:pointer; }
.upb-preset.active { color:white; border-color:var(--accent); background:var(--accent); }
.upb-context { display:flex; align-items:center; gap:5px; padding-bottom:7px; font-size:10px; }
.upb-body { min-height:0; overflow:auto; padding:14px; }
.upb-status,.upb-grid { width:min(100%,1300px); margin-right:auto; margin-left:auto; }
.upb-status { margin-bottom:10px; color:#465149; font-size:10px; }
.upb-grid { display:grid; grid-template-columns:repeat(2,minmax(0,640px)); justify-content:center; align-items:start; gap:14px; }
.upb-chart-card { box-sizing:border-box; width:100%; min-width:0; padding:10px; border:1px solid var(--line); border-radius:9px; background:white; }
.upb-chart-card.concordia { grid-column:1/-1; max-width:940px; justify-self:center; }
.upb-chart-card:not(.concordia) { max-width:640px; justify-self:center; }
.upb-grid.magmatic .upb-chart-card.concordia { grid-column:auto; max-width:640px; }
.upb-grid.magmatic .kde-card { display:none; }
.upb-chart-head { display:flex; flex-wrap:wrap; align-items:center; justify-content:space-between; gap:7px; margin-bottom:4px; font-size:11.5px; font-weight:800; }
.upb-zoom { display:flex; gap:4px; }
.upb-chart { min-height:330px; }
.upb-chart svg { display:block; width:100%; height:auto; touch-action:none; }
.upb-note { min-height:26px; margin-top:4px; color:var(--muted); font-size:9px; line-height:1.4; }
.upb-peaks { display:flex; flex-wrap:wrap; gap:4px 8px; margin-top:5px; color:#425048; font-size:9px; }
.point-card { padding:9px 10px; border:1px solid #cad6cf; border-left:3px solid var(--accent); border-radius:8px; background:#f7faf8; }
.point-card strong { display:block; margin-bottom:3px; font-size:11px; }
.point-card span { color:#4f5c55; font-size:10.5px; line-height:1.4; }
.popup-point { margin:-2px 0 7px; color:#68736c; font-size:10px; line-height:1.35; }
.chem-open { width:100%; margin-top:10px; padding:9px 11px; border:0; border-radius:7px; color:white; background:var(--accent); cursor:pointer; font-weight:700; font-size:11px; }
.chem-open:hover { background:#244f39; }
.chem-open:disabled { color:#89928c; background:#e3e7e4; cursor:not-allowed; }
.chem-modal { position:fixed; inset:0; z-index:45; display:none; background:#eef1ef; }
.chem-modal.visible { display:grid; grid-template-columns:230px minmax(0,1fr); }
.chem-sidebar { min-height:0; display:flex; flex-direction:column; border-right:1px solid var(--line); background:var(--panel); }
.chem-side-head { padding:16px; border-bottom:1px solid var(--line); }
.chem-side-head h2 { margin:4px 0; font-size:20px; }
.chem-actions { display:grid; grid-template-columns:1fr 1fr; gap:6px; margin-top:11px; }
.chem-actions button,.chem-close,.chem-sample { border:1px solid #cbd3cd; border-radius:7px; background:white; cursor:pointer; }
.chem-actions button { padding:7px 6px; font-size:10.5px; }
.chem-actions button:hover,.chem-close:hover,.chem-sample:hover { border-color:#5c9270; background:#edf5f0; }
.chem-ref-picker { margin-top:9px; border-top:1px solid var(--line); padding-top:9px; }
.chem-ref-picker summary { color:#465149; cursor:pointer; font-size:10.5px; font-weight:700; }
.chem-refs { display:grid; gap:5px; max-height:190px; margin-top:7px; overflow:auto; }
.chem-ref { display:grid; grid-template-columns:14px 1fr; gap:6px; align-items:start; color:#4f5c55; font-size:9px; line-height:1.25; cursor:pointer; }
.chem-ref input { width:auto; margin:1px 0 0; padding:0; }
.chem-samples { min-height:0; flex:1; overflow:auto; padding:8px; }
.chem-sample { display:flex; width:100%; align-items:center; gap:8px; margin-bottom:5px; padding:8px 9px; text-align:left; font-size:11px; }
.chem-sample.selected { border-color:#89a797; background:#edf5f0; }
.chem-sample.active { box-shadow:0 0 0 2px var(--accent); }
.chem-check { width:9px; height:9px; border-radius:2px; border:1px solid #96a29a; background:white; }
.chem-sample.selected .chem-check { border-color:var(--accent); background:var(--accent); }
.chem-group { margin-left:auto; color:var(--muted); font-size:8.5px; }
.chem-main { min-width:0; min-height:0; display:flex; flex-direction:column; }
.chem-toolbar { display:flex; align-items:center; gap:12px; padding:11px 16px; border-bottom:1px solid var(--line); background:white; }
.chem-toolbar-title { min-width:0; flex:1; }
.chem-toolbar-title strong { display:block; font-size:13px; }
.chem-toolbar-title span { color:var(--muted); font-size:10px; }
.chem-close { padding:8px 11px; font-size:11px; }
.chem-scroll { min-height:0; overflow:auto; padding:15px; }
.chem-section { margin-bottom:18px; }
.chem-section-head { display:flex; align-items:end; justify-content:space-between; gap:12px; margin-bottom:8px; }
.chem-section-head h3 { margin:0; font-size:14px; }
.chem-section-head span { color:var(--muted); font-size:9.5px; text-align:right; }
.chem-grid { display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:12px; }
.chart-card { min-width:0; padding:10px; border:1px solid var(--line); border-radius:9px; background:white; }
.chart-card.wide { grid-column:1/-1; }
.chart-card.tas-card { width:min(100%,760px); margin:0 auto; }
.chart-title { display:flex; justify-content:space-between; gap:8px; margin-bottom:3px; font-size:11.5px; font-weight:800; }
.chart-title span { color:var(--muted); font-size:9px; font-weight:400; }
.chart { width:100%; min-height:320px; }
.chart svg { display:block; width:100%; height:auto; overflow:visible; }
.chart-note { min-height:27px; margin-top:4px; color:var(--muted); font-size:9px; line-height:1.35; }
.chem-legend { display:flex; flex-wrap:wrap; gap:11px; margin:0 0 9px; color:#4f5c55; font-size:9.5px; }
.chem-legend-item { display:flex; align-items:center; gap:5px; }
.chem-swatch { width:16px; height:3px; background:#9ba59f; }
.chem-swatch.active { height:4px; background:#c53f31; }
.chem-swatch.morb { border-top:2px dashed #2c6da4; background:transparent; }
.chem-swatch.oib { border-top:2px dashed #b27718; background:transparent; }
.context-key { width:11px; height:11px; display:inline-block; flex:0 0 auto; border:2px solid currentColor; background:currentColor; }
.context-key.tuff { color:#287f73; border-radius:50%; }
.context-key.igarka_dyke { color:#7653a6; transform:rotate(45deg) scale(.78); }
.context-key.sill { color:#b06b2b; }
.context-key.other { color:#7b8580; clip-path:polygon(50% 0,100% 100%,0 100%); border:0; }
.chem-warning { margin-top:8px; padding:8px 10px; border-left:3px solid var(--prepared); color:#66552f; background:#faf5ea; font-size:9.5px; line-height:1.4; }
.lightbox { position:fixed; inset:0; z-index:50; display:none; grid-template-rows:auto minmax(0,1fr); padding:14px; background:rgba(8,13,10,.92); backdrop-filter:blur(5px); }
.lightbox.visible { display:grid; }
.lightbar { display:flex; justify-content:space-between; align-items:center; gap:12px; padding:2px 2px 12px; color:white; }
.light-title { min-width:0; font-size:12px; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.light-controls { display:flex; gap:6px; flex:0 0 auto; }
.light-controls button { min-width:36px; height:34px; padding:0 10px; border:1px solid rgba(255,255,255,.25); border-radius:7px; color:white; background:rgba(255,255,255,.12); cursor:pointer; }
.light-controls button:hover { background:rgba(255,255,255,.22); }
.lightstage { position:relative; min-height:0; display:grid; place-items:center; overflow:hidden; }
.lightstage img { width:100%; height:100%; min-width:0; min-height:0; display:block; object-fit:contain; filter:drop-shadow(0 12px 22px rgba(0,0,0,.38)); }
.lightnav { position:absolute; z-index:2; top:50%; width:46px; height:58px; transform:translateY(-50%); border:0; border-radius:8px; color:white; background:rgba(10,16,12,.48); font-size:28px; cursor:pointer; }
.lightnav:hover { background:rgba(10,16,12,.72); }
.lightprev { left:8px; }
.lightnext { right:8px; }
.list-empty { padding:18px 10px; color:var(--muted); font-size:12px; text-align:center; }
.workflow-modal { position:fixed; inset:0; display:none; grid-template-rows:auto minmax(0,1fr); color:var(--text); background:#f4f6f4; z-index:18; }
.workflow-modal.visible { display:grid; }
.workflow-toolbar { display:flex; flex-wrap:wrap; align-items:center; gap:8px; padding:11px 14px; border-bottom:1px solid var(--line); background:white; }
.workflow-toolbar strong { margin-right:auto; font-size:14px; }
.workflow-toolbar input,.workflow-toolbar select { width:auto; min-width:150px; }
.workflow-toolbar button { padding:8px 11px; border:1px solid #bdc8c0; border-radius:7px; color:var(--text); background:white; cursor:pointer; }
.workflow-body { min-height:0; overflow:auto; padding:14px; }
.workflow-explain { margin:0 0 10px; color:var(--muted); font-size:10.5px; line-height:1.45; }
.workflow-count { margin-bottom:7px; color:var(--muted); font-size:10px; }
.workflow-table-wrap { min-width:860px; overflow:hidden; border:1px solid var(--line); border-radius:9px; background:white; }
.workflow-table { width:100%; border-collapse:collapse; font-size:10.5px; }
.workflow-table th { position:sticky; top:0; z-index:1; padding:8px 7px; border-bottom:1px solid var(--line); color:#526059; background:#eef2ef; text-align:left; white-space:nowrap; }
.workflow-table td { padding:7px; border-bottom:1px solid #e6eae7; vertical-align:top; }
.workflow-table tr:last-child td { border-bottom:0; }
.workflow-table tbody tr { cursor:pointer; }
.workflow-table tbody tr:hover { background:#f3f7f4; }
.workflow-sample { font-weight:800; white-space:nowrap; }
.workflow-stage { max-width:180px; }
.workflow-detail { max-width:240px; color:var(--muted); font-size:9.5px; line-height:1.35; }
.workflow-scroll-x { overflow-x:auto; }
@media(max-width:1120px) { .app { grid-template-columns:270px 1fr 315px; } }
@media(max-width:980px) { .chem-modal.visible { grid-template-columns:190px minmax(0,1fr); } .chem-grid,.upb-grid,.geo-grid { grid-template-columns:1fr; } .chart-card.wide,.upb-chart-card.concordia { grid-column:auto; } .upb-chart-card.concordia { max-width:720px; } }
@media(max-width:860px) { html,body { overflow:auto; } .app { grid-template-columns:1fr; grid-template-rows:auto 62vh auto; height:auto; min-height:0; overflow:visible; } .panel { height:auto; } .map-wrap { height:62vh; } .left,.right { border:0; } .list { max-height:38vh; } .right { min-height:360px; } .chem-modal.visible { grid-template-columns:1fr; grid-template-rows:190px minmax(0,1fr); } .chem-sidebar { border-right:0; border-bottom:1px solid var(--line); } .chem-samples { display:flex; gap:5px; } .chem-sample { min-width:125px; margin:0; } }
</style>
</head>
<body>
<div class="app">
  <aside class="panel left">
    <div class="panel-head">
      <div class="eyebrow">Полевая коллекция</div>
      <h1>Игарка · образцы</h1>
      <div class="hint">Реестр 2025–2026. Выберите точку на карте или образец в списке.</div>
      <div id="summary" class="summary"></div>
    </div>
    <div class="controls">
      <input id="search" type="search" placeholder="Поиск образца, участка, породы">
      <div class="controls-row"><select id="year"></select><select id="status"></select></div>
    </div>
    <div id="listSummary" class="list-summary"></div>
    <div id="list" class="list"></div>
  </aside>
  <main class="map-wrap">
    <canvas id="map"></canvas>
    <div class="map-toolbar"><button id="zoomIn" title="Приблизить">+</button><button id="zoomOut" title="Отдалить">−</button><button id="fit" title="Показать отфильтрованные точки">Все точки</button><button id="literatureToggle" aria-pressed="true" title="Показать или скрыть точки Кочнева">Кочнев</button><button id="labelsToggle" aria-pressed="true" title="Показать или скрыть подписи точек">Подписи</button><button id="workflowOpen" title="Открыть таблицу состояния образцов">Состояние проб</button><button id="chemOpen" title="Открыть диаграммы">Геохимия</button><button id="geoOpen" title="Открыть KDE диаграммы">Геохронология</button></div>
    <div class="map-note">Спутниковая подложка Google</div>
    <div class="map-layer-control" aria-label="Слой геологической карты"><label class="map-layer-heading"><input id="mapOverlayToggle" type="checkbox" checked>Геологическая карта</label><div id="mapOverlaySlider" class="map-layer-slider"><label for="mapOverlayOpacity">Прозрачность <output id="mapOverlayValue" for="mapOverlayOpacity">35%</output></label><input id="mapOverlayOpacity" type="range" min="0" max="100" step="1" value="35" aria-label="Прозрачность геологической карты"></div><div id="mapOverlayStatus" class="map-layer-status" role="status"></div></div>
    <div class="map-legend"><span class="map-legend-item"><i class="legend-pin y2025"></i>пробы 2025</span><span class="map-legend-item"><i class="legend-pin y2026"></i>пробы 2026</span><span class="map-legend-item" id="kochnevLegend"><i class="legend-pin kochnev"></i>точки Кочнева</span></div>
    <div id="popup" class="popup"></div>
  </main>
  <section class="panel right">
    <div class="panel-head"><div id="detailEyebrow" class="eyebrow">Карточка образца</div><h1 id="detailTitle">Выберите образец</h1><div id="detailSub" class="hint"></div></div>
    <div id="detail" class="detail-scroll"></div>
  </section>
</div>
<div id="workflowModal" class="workflow-modal" aria-hidden="true">
  <div class="workflow-toolbar">
    <strong>Состояние образцов и запланированных анализов</strong>
    <input id="workflowSearch" type="search" placeholder="Поиск образца или объекта">
    <select id="workflowYear"></select>
    <select id="workflowFilter"></select>
    <button id="workflowClose">Закрыть</button>
  </div>
  <div class="workflow-body">
    <p class="workflow-explain">В таблице показаны только работы, относящиеся к плану конкретного образца. Пустая ячейка реестра означает «не запланировано», значение «нет» — работа запланирована, но пока не выполнена. Явные указания в описании, статусе и комментарии также учитываются.</p>
    <div id="workflowCount" class="workflow-count"></div>
    <div class="workflow-scroll-x"><div class="workflow-table-wrap"><table class="workflow-table"><thead><tr><th>Образец</th><th>Объект</th><th>Текущий этап</th><th>Шлиф</th><th>Химсостав</th><th>U–Pb</th><th>Sm–Nd</th><th>Треки AFT</th><th>Основание статуса</th></tr></thead><tbody id="workflowRows"></tbody></table></div></div>
  </div>
</div>
<div id="chemModal" class="chem-modal" aria-hidden="true">
  <aside class="chem-sidebar">
    <div class="chem-side-head">
      <div class="eyebrow">Интерактивные диаграммы</div>
      <h2>Геохимия</h2>
      <div id="chemCount" class="hint"></div>
      <div class="chem-actions"><button id="chemCurrent">Текущий</button><button id="chemAll">Все пробы</button></div>
      <details class="chem-ref-picker"><summary>Референтные линии</summary><div id="chemRefs" class="chem-refs"></div></details>
    </div>
    <div id="chemSamples" class="chem-samples"></div>
  </aside>
  <main class="chem-main">
    <div class="chem-toolbar"><div class="chem-toolbar-title"><strong id="chemActiveTitle">Выбранный образец</strong><span>Щёлкните точку на диаграмме, чтобы сделать образец активным</span></div><button id="chemClose" class="chem-close">Закрыть</button></div>
    <div class="chem-scroll">
      <div class="chem-legend"><span class="chem-legend-item"><i class="chem-swatch active"></i>активный</span><span class="chem-legend-item"><i class="context-key tuff"></i>туфы</span><span class="chem-legend-item"><i class="context-key igarka_dyke"></i>дайки Игарской толщи</span><span class="chem-legend-item"><i class="context-key sill"></i>силлы</span><span class="chem-legend-item"><i class="context-key other"></i>прочие</span></div>
      <section class="chem-section"><div class="chem-section-head"><h3>1. TAS · классификация вулканических пород</h3><span>Le Bas et al., 1986 · безводный пересчёт</span></div><div class="chem-grid"><article class="chart-card wide tas-card"><div class="chart-title">SiO₂ — Na₂O + K₂O <span>мас. %</span></div><div id="tasChart" class="chart"></div><div id="tasNote" class="chart-note"></div></article></div></section>
      <section class="chem-section"><div class="chem-section-head"><h3>2. Нормированные спектры</h3><span>Sun &amp; McDonough, 1989</span></div><div class="chem-grid"><article class="chart-card"><div class="chart-title">Редкоземельные элементы <span>C1-хондрит</span></div><div id="reeChart" class="chart"></div><div class="chart-note">Включённые в боковой панели референты нормированы на тот же C1-хондрит.</div></article><article class="chart-card"><div class="chart-title">Мультиэлементная диаграмма <span>примитивная мантия</span></div><div id="spiderChart" class="chart"></div><div class="chart-note">Состав и порядок элементов взяты из PM_1989.csv. K, P и Ti пересчитаны из K₂O, P₂O₅ и TiO₂.</div></article></div></section>
      <section class="chem-section"><div class="chem-section-head"><h3>3. Диаграммы по малоподвижным элементам</h3><span>логарифмические шкалы</span></div><div class="chem-grid"><article class="chart-card"><div class="chart-title">TiO₂/Yb — Th/Nb <span>Pearce et al., 2021</span></div><div id="lipChart" class="chart"></div><div class="chart-note">LIP printing: TiO₂ в мас. %, Yb, Th и Nb в ppm. Полые знаки не прошли базитовый фильтр Pearce (1996), использованный авторами статьи.</div></article><article class="chart-card"><div class="chart-title">Th/Yb — Ta/Yb <span>Pearce, 1982</span></div><div id="thChart" class="chart"></div><div class="chart-note">Серая полоса — мантийный массив между Th/Ta = 1 и 2; смещение вверх отражает относительное обогащение Th.</div></article></div></section>
      <div class="chem-warning">TAS чувствителен к вторичным изменениям и высоким ППП, поэтому построен после пересчёта главных оксидов без ППП. Для изменённых пород диаграммы по малоподвижным элементам обычно устойчивее. Приблизительные ICP-MS значения V, Cr, Co, Ni, Cu, Zn, Mo, Sn и Tl в этих пяти графиках не используются.</div>
    </div>
  </main>
</div>
<div id="geoModal" class="geo-modal" aria-hidden="true">
  <div class="geo-toolbar">
    <div class="geo-heading"><strong>Геохронология · KDE</strong><span>Щёлкните пик на графике или кнопку под ним, чтобы рассчитать взвешенное среднее</span></div>
    <label class="geo-control">Ширина ядра <span id="geoBandwidthLabel">30 млн лет</span><input id="geoBandwidth" type="range" min="5" max="100" step="5" value="30"></label>
    <button id="geoClose" class="geo-close">Закрыть</button>
  </div>
  <div class="geo-scroll">
    <div class="geo-note">Выборка пика ограничивается соседними локальными минимумами KDE. Границы и среднее меняются вместе с шириной ядра; расчёт не заменяет concordia age или формальный MDA.</div>
    <div id="geoGrid" class="geo-grid"></div>
  </div>
</div>
<div id="youngModal" class="young-modal" aria-hidden="true">
  <div class="young-toolbar">
    <strong>Губинская свита · анализ молодых пиков</strong>
    <div id="youngSeriesButtons" class="young-series-buttons"></div>
    <button id="youngTheme">Светлая тема</button>
    <button id="youngClose">Закрыть</button>
  </div>
  <div class="young-content">
    <section class="young-panel">
      <div id="youngMeta" class="young-meta"></div>
      <div id="youngChart" class="young-chart"></div>
      <div class="young-note">Области 700–750, 750–850 и 850–900 млн лет заданы как в исходной заготовке. Проценты рассчитаны от числа принятых анализов внутри диапазона 700–900 млн лет.</div>
    </section>
  </div>
</div>
<div id="upbModal" class="upb-modal" aria-hidden="true">
  <div class="upb-toolbar">
    <div class="upb-heading"><strong id="upbTitle">U–Pb</strong><span id="upbSubtitle">Интерактивный пересчёт выборки</span></div>
    <label class="upb-control">Метод дискордантности<select id="upbMethod"><option value="hybrid">Гибридный D1/D2 (как сейчас)</option><option value="d1">D1: возраст ²⁰⁷Pb/²³⁵U vs ²⁰⁶Pb/²³⁸U</option><option value="d2">D2: возраст ²⁰⁷Pb/²⁰⁶Pb vs ²⁰⁶Pb/²³⁸U</option></select></label>
    <label class="upb-control">Порог: <b id="upbLimitLabel">10%</b><span class="upb-slider-row"><input id="upbLimit" type="range" min="0" max="20" step="0.5" value="10"><button class="upb-preset" data-limit="5">5%</button><button class="upb-preset active" data-limit="10">10%</button></span></label>
    <label class="upb-context"><input id="upbContext" type="checkbox" checked> исключать контекстно отмеченные анализы</label>
    <button id="upbCompare" class="upb-close" style="display:none">Сравнить 1812-1 / SM25-07</button>
    <button id="upbClose" class="upb-close">Закрыть</button>
  </div>
  <div class="upb-body">
    <div id="upbStatus" class="upb-status"></div>
    <div id="upbGrid" class="upb-grid">
      <section class="upb-chart-card concordia"><div class="upb-chart-head"><span>Wetherill concordia ²⁰⁷Pb/²³⁵U — ²⁰⁶Pb/²³⁸U</span><span class="upb-zoom"><button id="upbZoomIn">+</button><button id="upbZoomOut">−</button><button id="upbZoomFit">По данным</button><button id="upbZoomReset">Вся</button></span></div><div id="upbConcordia" class="upb-chart"></div><div class="upb-note">Колесо или щипок — масштаб, перетаскивание — панорама. Двойной щелчок приближает. Эллипсы 2σ по исходным ошибкам 1σ и Rho3.</div></section>
      <section class="upb-chart-card kde-card"><div class="upb-chart-head"><span>KDE и гистограмма</span><span>h = 30 млн лет</span></div><div id="upbKde" class="upb-chart"></div><div id="upbPeaks" class="upb-peaks"></div></section>
      <section class="upb-chart-card"><div class="upb-chart-head"><span>Weighted mean</span><span>возраст по горизонтали</span></div><div id="upbWm" class="upb-chart"></div><div id="upbWmNote" class="upb-note"></div></section>
    </div>
  </div>
</div>
<div id="lightbox" class="lightbox" aria-hidden="true">
  <div class="lightbar"><div id="lightTitle" class="light-title"></div><div class="light-controls"><button id="lightClose">Закрыть</button></div></div>
  <div class="lightstage"><button id="lightPrev" class="lightnav lightprev" title="Предыдущее фото">‹</button><img id="lightImage" alt=""><button id="lightNext" class="lightnav lightnext" title="Следующее фото">›</button></div>
</div>
<script>
const samples = __PAYLOAD__;
const GEOCHRONOLOGY = __GEOCHRONOLOGY__;
const literaturePoints = __LITERATURE_POINTS__;
const RASTER_OVERLAY = __RASTER_OVERLAY__;
const STATUS = {
  completed:{label:'План выполнен',color:'#23865b',note:'Все анализы, запланированные для этого образца, выполнены'},
  lab:{label:'В лаборатории',color:'#3278bd',note:'Одна или несколько запланированных работ выполняются в лаборатории'},
  preparation:{label:'Подготовка материала',color:'#b97817',note:'Идёт выделение материала или материал уже подготовлен'},
  planned:{label:'Запланировано',color:'#6c7470',note:'Работы указаны в плане, но ещё не начаты'},
  no_plan:{label:'План не указан',color:'#8b948f',note:'В реестре не указаны анализы для этого образца'}
};
const TASK_STATE={done:'готово',in_lab:'в лаборатории',preparing:'подготовка',prepared:'материал готов',planned:'план',not_planned:'—'};
const C1={La:.237,Ce:.612,Pr:.095,Nd:.467,Sm:.153,Eu:.058,Gd:.2055,Tb:.0374,Dy:.254,Ho:.0566,Er:.1655,Tm:.0255,Yb:.17,Lu:.0254};
const PM=__PRIMITIVE_MANTLE__;
const CHEM_REFERENCES=__REFERENCES__;
const REE_ELEMENTS=['La','Ce','Pr','Nd','Sm','Eu','Gd','Tb','Dy','Ho','Er','Tm','Yb','Lu'];
const SPIDER_ELEMENTS=__SPIDER_ELEMENTS__;
const CONTEXT={tuff:{label:'Туфы',color:'#287f73',shape:'circle'},igarka_dyke:{label:'Дайки Игарской толщи',color:'#7653a6',shape:'diamond'},sill:{label:'Силлы',color:'#b06b2b',shape:'square'},other:{label:'Прочие',color:'#7b8580',shape:'triangle'}};
const TAS_FIELDS=[
 {id:'Pc',label:'Пикрит',p:[[41,3],[45,3],[45,2],[45,0],[41,0]]},
 {id:'Ba',label:'Щел. базальт',p:[[45,5],[52,5],[45,2]]},
 {id:'Bs',label:'Субщел. базальт',p:[[45,2],[52,5],[52,0],[45,0]]},
 {id:'F',label:'Фоидит',p:[[35,9],[37,14],[52.5,18],[52.5,14],[48.4,11.5],[45,9.4],[41,7],[41,3],[37,3]]},
 {id:'O1',label:'Баз. андезит',p:[[52,0],[52,5],[57,5.9],[57,0]]},
 {id:'O2',label:'Андезит',p:[[57,0],[57,5.9],[63,7],[63,0]]},
 {id:'O3',label:'Дацит',p:[[63,0],[63,7],[69,8],[77.3,0]]},
 {id:'Ph',label:'Фонолит',p:[[52.5,14],[52.5,18],[57,18],[63,16.2],[61,13.5],[57.6,11.7]]},
 {id:'R',label:'Риолит',p:[[69,8],[69,13],[85.9,6.8],[87.5,4.7],[77.3,0]]},
 {id:'S1',label:'Трахибазальт',p:[[45,5],[49.4,7.3],[52,5]]},
 {id:'S2',label:'Баз. трахиандезит',p:[[49.4,7.3],[53,9.3],[57,5.9],[52,5]]},
 {id:'S3',label:'Трахиандезит',p:[[53,9.3],[57.6,11.7],[61,8.6],[63,7],[57,5.9]]},
 {id:'T1',label:'Трахит',p:[[57.6,11.7],[61,13.5],[63,16.2],[69,13],[61,8.6]]},
 {id:'T2',label:'Трахидацит',p:[[61,8.6],[69,13],[69,8],[63,7]]},
 {id:'U1',label:'Тефрит',p:[[41,3],[41,7],[45,9.4],[49.4,7.3],[45,5],[45,3]]},
 {id:'U2',label:'Фонотефрит',p:[[45,9.4],[48.4,11.5],[53,9.3],[49.4,7.3]]},
 {id:'U3',label:'Тефрифонолит',p:[[48.4,11.5],[52.5,14],[57.6,11.7],[53,9.3]]}
];
let chemSelected=new Set(),chemActive=null,chemRefSelected=new Set();
const canvas=document.getElementById('map'), ctx=canvas.getContext('2d'), popup=document.getElementById('popup');
const tileCache=new Map(), hitItems=[];
let dpr=window.devicePixelRatio||1, zoom=9, targetZoom=9, zoomAnimation=0, center={lat:67.45,lon:86.55}, dragging=false, moved=false, dragStart=null, activeSample=null, activeLiterature=null, showLiteraturePoints=!document.documentElement.classList.contains('mapshot'), showMapLabels=true;
const mapshotSamples=new Set((new URLSearchParams(location.search).get('samples')||'').split(',').map(v=>v.trim()).filter(Boolean));
const tileTemplate='https://mt1.google.com/vt/lyrs=s&x={x}&y={y}&z={z}';
function esc(v){return String(v??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));}
function statusLabel(key){return STATUS[key]?.label||STATUS.no_plan.label;}
function norm(v){return String(v||'').toLowerCase().trim();}
function tasks(s){return s.workflow_tasks||[];}
function taskBy(s,key){return tasks(s).find(t=>t.key===key)||{state:'not_planned',detail:'Не запланировано',label:key};}
function filtered(){const q=norm(document.getElementById('search').value),year=document.getElementById('year').value,status=document.getElementById('status').value;return samples.filter(s=>(!mapshotSamples.size||mapshotSamples.has(s.sample))&&(year==='all'||String(s.year)===year)&&(status==='all'||s.workflow_status===status)&&(!q||[s.sample,s.location,s.description,s.comment,s.status_raw,...tasks(s).map(t=>t.detail)].join(' ').toLowerCase().includes(q)));}
function initControls(){const years=['all',...new Set(samples.map(s=>String(s.year)))];document.getElementById('year').innerHTML=years.map(v=>`<option value="${v}">${v==='all'?'Все годы':v}</option>`).join('');document.getElementById('status').innerHTML='<option value="all">Все статусы</option>'+Object.entries(STATUS).map(([k,v])=>`<option value="${k}">${v.label}</option>`).join('');}
function renderSummary(){const counts={completed:0,lab:0,preparation:0,planned:0,no_plan:0};samples.forEach(s=>counts[s.workflow_status]++);document.getElementById('summary').innerHTML=[['completed','выполнен'],['lab','в лаб.'],['preparation','подгот.'],['planned','план'],['no_plan','без плана']].map(([k,label])=>`<div class="metric"><strong style="color:${STATUS[k].color}">${counts[k]}</strong><span>${label}</span></div>`).join('');}
function renderList(){const arr=filtered();document.getElementById('listSummary').textContent=`Показано ${arr.length} из ${samples.length}`;document.getElementById('list').innerHTML=arr.map(s=>{const chips=tasks(s).filter(t=>t.state!=='not_planned').map(t=>`<span class="task-chip ${t.state}" title="${esc(t.label)}: ${esc(t.detail)}">${esc(t.label)} · ${TASK_STATE[t.state]}</span>`).join('');return `<div class="item ${s.sample===activeSample?'active':''}" data-sample="${esc(s.sample)}"><div class="item-top"><span class="sample">${esc(s.sample)}</span><span class="year">${s.year}</span></div><div class="item-location">${esc(s.location||s.description||'Без описания')}</div><div class="status-line"><span class="status-name"><b>${statusLabel(s.workflow_status)}</b></span></div><div class="task-chips">${chips||'<span class="task-chip planned">анализы не указаны</span>'}</div></div>`;}).join('')||'<div class="list-empty">По фильтрам ничего не найдено</div>';document.querySelectorAll('.item').forEach(el=>el.onclick=()=>openSample(el.dataset.sample,true));draw();}
function workflowCell(task){return `<span class="task-state ${task.state}" title="${esc(task.detail)}">${TASK_STATE[task.state]}</span>`;}
function workflowFiltered(){const q=norm(document.getElementById('workflowSearch').value),year=document.getElementById('workflowYear').value,status=document.getElementById('workflowFilter').value;return samples.filter(s=>(year==='all'||String(s.year)===year)&&(status==='all'||s.workflow_status===status)&&(!q||[s.sample,s.location,s.description,s.comment,...tasks(s).map(t=>t.detail)].join(' ').toLowerCase().includes(q)));}
function renderWorkflowTable(){const arr=workflowFiltered(),body=document.getElementById('workflowRows');document.getElementById('workflowCount').textContent=`Показано ${arr.length} из ${samples.length}`;body.innerHTML=arr.map(s=>`<tr data-sample="${esc(s.sample)}"><td class="workflow-sample">${esc(s.sample)}</td><td>${esc(s.location||s.description||'—')}</td><td class="workflow-stage"><b>${statusLabel(s.workflow_status)}</b><br><span class="workflow-detail">${esc(STATUS[s.workflow_status].note)}</span></td>${['thin_section','chemistry','upb','sm_nd','aft'].map(k=>`<td>${workflowCell(taskBy(s,k))}</td>`).join('')}<td class="workflow-detail">${esc([s.status_raw,s.comment].filter(Boolean).join(' · ')||'—')}</td></tr>`).join('')||'<tr><td colspan="9">По фильтрам ничего не найдено</td></tr>';body.querySelectorAll('tr[data-sample]').forEach(row=>row.onclick=()=>{closeWorkflow();openSample(row.dataset.sample,true);});}
function openWorkflow(){const modal=document.getElementById('workflowModal'),years=['all',...new Set(samples.map(s=>String(s.year)))];document.getElementById('workflowYear').innerHTML=years.map(v=>`<option value="${v}">${v==='all'?'Все годы':v}</option>`).join('');document.getElementById('workflowFilter').innerHTML='<option value="all">Все этапы</option>'+Object.entries(STATUS).map(([k,v])=>`<option value="${k}">${v.label}</option>`).join('');document.getElementById('workflowSearch').value='';modal.classList.add('visible');modal.setAttribute('aria-hidden','false');renderWorkflowTable();}
function closeWorkflow(){const modal=document.getElementById('workflowModal');modal.classList.remove('visible');modal.setAttribute('aria-hidden','true');}
function field(label,value){return value?`<div class="field"><b>${esc(label)}</b>${esc(value)}</div>`:'';}
function photoGallery(s){if(!s.photos?.length)return '<div class="photo-empty">Для этой точки фотографии не найдены</div>';return `<div class="photo-grid">${s.photos.map((p,i)=>`<button class="photo-thumb-button outcrop-photo" data-photo-index="${i}" title="${esc(p.name)}"><img src="${esc(p.thumb)}" alt="${esc(s.photo_point)} · фото ${i+1}" loading="lazy"><span class="photo-number">${i+1}</span></button>`).join('')}</div>`;}
function thinSectionGallery(s){if(!s.thin_section_photos?.length)return '<div class="photo-empty">Для этого образца фотографии шлифов не найдены</div>';return `<div class="photo-grid">${s.thin_section_photos.map((p,i)=>`<button class="photo-thumb-button thin-section-photo" data-photo-index="${i}" title="${esc(p.name)}"><img src="${esc(p.thumb)}" alt="${esc(s.sample)} · ${esc(p.mode)}" loading="lazy"><span class="photo-mode">${esc(p.mode)}</span><span class="photo-number">${i+1}</span></button>`).join('')}</div>`;}
function petrographyHtml(s){if(!s.petrography?.length)return '<div class="photo-empty">Петрографическое описание отсутствует</div>';return `<div class="petro-list">${s.petrography.map(p=>`<article class="petro-card"><div class="petro-heading"><span>${esc(p.section)}</span>${esc(p.title)}</div><div class="petro-text">${esc(p.description)}</div>${p.note?`<div class="petro-note">${esc(p.note)}</div>`:''}<div class="petro-source">Источник: shlify_igarka_opisaniya.md</div></article>`).join('')}</div>`;}
function upbHtml(s){if(!s.upb)return '';const compare=s.sample==='SM25-07'?'<button class="chem-open upb-compare-open">Сравнить с 1812-1</button>':'';return `<div class="section"><div class="section-title">U–Pb датирование</div><div class="upb-card"><button class="upb-figure" title="Открыть интерактивные U–Pb диаграммы"><img src="${esc(s.upb.figure)}" alt="U–Pb результаты ${esc(s.sample)}" loading="lazy"></button><div class="upb-summary">${esc(s.upb.summary)}${s.upb.interpretation?`<div class="upb-interpretation">${esc(s.upb.interpretation)}</div>`:''}<div class="upb-links"><button class="chem-open sample-upb-open">Открыть интерактивно</button>${compare}<a href="${esc(s.upb.figure)}" target="_blank">Статический рисунок</a><a href="${esc(s.upb.audit)}" download>Таблица аудита CSV</a><a href="upb/UPb_audit_report.md" target="_blank">Методический аудит</a></div></div></div></div>`;}
function literatureUpbHtml(p){const figure=p.figure?`<button class="upb-figure literature-upb-open" data-dataset="${esc(p.dataset_id)}" title="Открыть интерактивные U–Pb диаграммы"><img src="${esc(p.figure)}" alt="U–Pb результаты ${esc(p.sample)}" loading="lazy"></button>`:'',primary=p.figure?`<button class="chem-open literature-upb-open" data-dataset="${esc(p.dataset_id)}">Открыть интерактивно</button><button class="chem-open literature-compare-open">Сравнить с SM25-07</button>`:`<button class="chem-open literature-geo-open" data-dataset="${esc(p.dataset_id)}">Открыть KDE</button>`;return `<div class="section"><div class="section-title">U–Pb датирование</div><div class="upb-card">${figure}<div class="upb-summary"><b>${esc(p.summary)}</b><div class="upb-interpretation">${esc(p.detail)}</div><div class="upb-links">${primary}<a href="${esc(p.audit)}" download>Таблица аудита CSV</a><a href="${esc(p.doi)}" target="_blank" rel="noopener">Статья / DOI</a></div></div></div></div>`;}
function literatureGeochHtml(s){const c=s.literature_geochronology;if(!c)return '';return `<div class="section"><div class="section-title">Литературные данные по геохронологии</div><article class="literature-card"><strong>${esc(c.title)}</strong><p>${esc(c.summary)}</p><p>${esc(c.detail)}</p><div class="literature-source">Источник: Kochnev et al., 2022${c.sample?` · образец ${esc(c.sample)}`:''}${c.analytical_prefix?` · аналитический префикс ${esc(c.analytical_prefix)}`:''}</div><div class="literature-actions">${c.dataset_id?`<button class="literature-geo-open" data-dataset="${esc(c.dataset_id)}">${esc(c.button_label||'Открыть KDE Кочнева')}</button>`:''}${c.audit?`<a href="${esc(c.audit)}" download>Аудит CSV</a>`:''}<a href="${esc(c.doi)}" target="_blank" rel="noopener">Статья / DOI</a></div></article></div>`;}
function pointHtml(s){if(!s.point_description)return '';return `<div class="section"><div class="section-title">Точка / обнажение</div><div class="point-card"><strong>${esc(s.point_id)}</strong><span>${esc(s.point_description)}${s.point_note?`<br>${esc(s.point_note)}`:''}</span></div></div>`;}
function chemistryHtml(s){return `<div class="section"><div class="section-title">Геохимические диаграммы</div><button class="chem-open sample-chem-open" ${s.geochem?'':'disabled'}>${s.geochem?'Открыть геохимию образца':'Для образца нет химических данных'}</button></div>`;}
let lightboxPhotos=[],lightboxIndex=0,lightboxPoint='',lightboxKind='фото';
function renderLightbox(){const p=lightboxPhotos[lightboxIndex];if(!p)return;document.getElementById('lightImage').src=p.full;document.getElementById('lightImage').alt=`${lightboxPoint} · ${lightboxKind} ${lightboxIndex+1}`;document.getElementById('lightTitle').textContent=`${lightboxPoint} · ${lightboxKind} ${lightboxIndex+1} из ${lightboxPhotos.length} · ${p.name}`;const multiple=lightboxPhotos.length>1;document.getElementById('lightPrev').style.display=multiple?'block':'none';document.getElementById('lightNext').style.display=multiple?'block':'none';}
function openLightbox(items,point,index,kind='фото'){lightboxPhotos=items||[];lightboxPoint=point;lightboxIndex=index;lightboxKind=kind;renderLightbox();const box=document.getElementById('lightbox');box.classList.add('visible');box.setAttribute('aria-hidden','false');}
function closeLightbox(){const box=document.getElementById('lightbox');box.classList.remove('visible');box.setAttribute('aria-hidden','true');document.getElementById('lightImage').removeAttribute('src');}
function stepLightbox(delta){if(!lightboxPhotos.length)return;lightboxIndex=(lightboxIndex+delta+lightboxPhotos.length)%lightboxPhotos.length;renderLightbox();}
function openSample(sampleId,pan){const s=samples.find(x=>x.sample===sampleId);if(!s)return;activeSample=sampleId;activeLiterature=null;document.getElementById('detailEyebrow').textContent='Карточка образца';document.getElementById('detailTitle').textContent=s.sample;document.getElementById('detailSub').textContent=`${s.year} · ${s.location||'участок не указан'}`;const st=STATUS[s.workflow_status];const trackHtml=tasks(s).map(t=>`<div class="track"><span class="track-name">${esc(t.label)}</span><span class="task-state ${t.state}">${TASK_STATE[t.state]}</span><span>${esc(t.detail)}</span></div>`).join('');const coords=Number.isFinite(s.lat)&&Number.isFinite(s.lon)?`${s.lat.toFixed(6)}, ${s.lon.toFixed(6)}${s.coords_inherited?' · координаты точки '+s.point_id:''}`:'нет координат';document.getElementById('detail').innerHTML=`<div class="hero-status"><div><strong>${st.label}</strong><span>${st.note}</span></div></div>${pointHtml(s)}${chemistryHtml(s)}${upbHtml(s)}${literatureGeochHtml(s)}<div class="section"><div class="section-title">Фото обнажения ${esc(s.photo_point)} · ${s.photos.length}</div>${photoGallery(s)}</div><div class="section"><div class="section-title">Фото шлифов ${esc(s.sample)} · ${s.thin_section_photos.length}</div>${thinSectionGallery(s)}</div><div class="section"><div class="section-title">Петрографическое описание · ${s.petrography.length}</div>${petrographyHtml(s)}</div><div class="section"><div class="section-title">План и состояние анализов</div><div class="tracks">${trackHtml}</div></div><div class="section"><div class="section-title">Описание образца</div>${field('Порода / материал',s.location)}${field('Полевое описание',s.description)}${field('Исходный статус',s.status_raw)}${field('Комментарий',s.comment)}</div><div class="section"><div class="section-title">Положение</div><div class="field coord"><b>Координаты WGS 84</b>${coords}</div></div><div class="source-note">Источник образца: лист Sheet2 файла «игарка_реестр.xlsx», строка ${s.row}.${s.point_row!==s.row?` Описание точки: строка ${s.point_row}.`:''} Пустая ячейка анализа трактуется как «не запланировано», значение «нет» — как запланированная, но не выполненная работа. Исходный реестр не изменён.</div>`;document.getElementById('detail').scrollTop=0;document.querySelectorAll('.outcrop-photo').forEach(button=>button.onclick=()=>openLightbox(s.photos,s.photo_point,Number(button.dataset.photoIndex)));document.querySelectorAll('.thin-section-photo').forEach(button=>button.onclick=()=>openLightbox(s.thin_section_photos,s.sample,Number(button.dataset.photoIndex)));document.querySelectorAll('.upb-figure,.sample-upb-open').forEach(button=>button.onclick=()=>openUpb(s.sample));document.querySelectorAll('.upb-compare-open').forEach(button=>button.onclick=openUpbComparison);document.querySelectorAll('.sample-chem-open').forEach(button=>button.onclick=()=>openChemistry(s.sample));document.querySelectorAll('.literature-geo-open').forEach(button=>button.onclick=()=>openGeochronology(button.dataset.dataset));document.querySelectorAll('.item').forEach(el=>el.classList.toggle('active',el.dataset.sample===sampleId));if(pan&&Number.isFinite(s.lat)&&Number.isFinite(s.lon)){center={lat:s.lat,lon:s.lon};zoom=Math.max(zoom,13);popup.classList.remove('visible');}draw();}
function openLiterature(pointId,pan){const p=literaturePoints.find(x=>x.id===pointId);if(!p)return;activeLiterature=pointId;activeSample=null;document.getElementById('detailEyebrow').textContent='Точка Кочнева';document.getElementById('detailTitle').textContent=p.sample;document.getElementById('detailSub').textContent=p.formation;document.getElementById('detail').innerHTML=`<div class="section"><div class="section-title">Образец и объект</div>${field('Порода / материал',p.material)}${field('Стратиграфия / положение',p.formation)}</div>${literatureUpbHtml(p)}<div class="section"><div class="section-title">Положение</div><div class="field coord"><b>Координаты WGS 84</b>${p.lat.toFixed(6)}, ${p.lon.toFixed(6)}</div>${field('Точность положения',p.coordinate_note)}</div><div class="source-note">Источник: ${esc(p.author)} · <a href="${esc(p.doi)}" target="_blank" rel="noopener">статья / DOI</a>. Точка Кочнева не входит в реестр образцов Малышева и не участвует в таблице их готовности.</div>`;document.getElementById('detail').scrollTop=0;document.querySelectorAll('.literature-geo-open').forEach(button=>button.onclick=()=>openGeochronology(button.dataset.dataset));document.querySelectorAll('.literature-upb-open').forEach(button=>button.onclick=()=>openLiteratureUpb(button.dataset.dataset));document.querySelectorAll('.literature-compare-open').forEach(button=>button.onclick=openUpbComparison);document.querySelectorAll('.item').forEach(el=>el.classList.remove('active'));if(pan){center={lat:p.lat,lon:p.lon};zoom=Math.max(zoom,13);popup.classList.remove('visible');}draw();}
const SVG_NS='http://www.w3.org/2000/svg';
function svgNode(tag,attrs={},content=''){const n=document.createElementNS(SVG_NS,tag);Object.entries(attrs).forEach(([k,v])=>n.setAttribute(k,v));if(content)n.textContent=content;return n;}
function chartFrame(container,{xDomain,yDomain,logX=false,logY=false,xTicks=[],yTicks=[],xLabel='',yLabel='',height=330}){
  container.innerHTML='';const width=Math.max(380,container.clientWidth||600),m={l:58,r:18,t:16,b:47},pw=width-m.l-m.r,ph=height-m.t-m.b;
  const svg=svgNode('svg',{viewBox:`0 0 ${width} ${height}`,role:'img','aria-label':`${yLabel} от ${xLabel}`});container.appendChild(svg);
  const scale=(v,d,a,b,log)=>{const q=log?Math.log10(v):v,q0=log?Math.log10(d[0]):d[0],q1=log?Math.log10(d[1]):d[1];return a+(q-q0)/(q1-q0)*(b-a);};
  const x=v=>scale(v,xDomain,m.l,m.l+pw,logX),y=v=>scale(v,yDomain,m.t+ph,m.t,logY);
  const grid=svgNode('g');svg.appendChild(grid);
  yTicks.forEach(t=>{const py=y(t);grid.appendChild(svgNode('line',{x1:m.l,y1:py,x2:m.l+pw,y2:py,stroke:'#dfe4e0','stroke-width':'1'}));grid.appendChild(svgNode('text',{x:m.l-8,y:py+3,'text-anchor':'end',fill:'#68736c','font-size':'9'},String(t)));});
  xTicks.forEach(t=>{const px=x(t.value??t);grid.appendChild(svgNode('line',{x1:px,y1:m.t,x2:px,y2:m.t+ph,stroke:'#eef1ef','stroke-width':'1'}));grid.appendChild(svgNode('text',{x:px,y:m.t+ph+16,'text-anchor':'middle',fill:'#68736c','font-size':'9'},String(t.label??t)));});
  svg.appendChild(svgNode('line',{x1:m.l,y1:m.t+ph,x2:m.l+pw,y2:m.t+ph,stroke:'#59645d'}));svg.appendChild(svgNode('line',{x1:m.l,y1:m.t,x2:m.l,y2:m.t+ph,stroke:'#59645d'}));
  svg.appendChild(svgNode('text',{x:m.l+pw/2,y:height-7,'text-anchor':'middle',fill:'#465149','font-size':'10'},xLabel));
  const yl=svgNode('text',{x:12,y:m.t+ph/2,'text-anchor':'middle',fill:'#465149','font-size':'10',transform:`rotate(-90 12 ${m.t+ph/2})`},yLabel);svg.appendChild(yl);
  return{svg,x,y,m,pw,ph,width,height};
}
function currentChemSamples(){return samples.filter(s=>s.geochem&&chemSelected.has(s.sample));}
function anhydrousMajor(s){const major=s.geochem.major,keys=['Na2O','MgO','Al2O3','SiO2','P2O5','K2O','CaO','TiO2','MnO','Fe2O3'],sum=keys.reduce((a,k)=>a+(major[k]||0),0),factor=sum?100/sum:1;return{factor,sum,SiO2:(major.SiO2||0)*factor,alkali:((major.Na2O||0)+(major.K2O||0))*factor};}
function pointInPoly(x,y,poly){let inside=false;for(let i=0,j=poly.length-1;i<poly.length;j=i++){const xi=poly[i][0],yi=poly[i][1],xj=poly[j][0],yj=poly[j][1];if(((yi>y)!=(yj>y))&&(x<(xj-xi)*(y-yi)/(yj-yi)+xi))inside=!inside;}return inside;}
function activateChem(sampleId){chemActive=sampleId;chemSelected.add(sampleId);renderChemSidebar();renderChemCharts();}
function plotPoint(frame,s,xv,yv,{hollow=false}={}){if(!(xv>0||!frame.logX)||!(yv>0||!frame.logY))return;const active=s.sample===chemActive,style=CONTEXT[s.geochem.context]||CONTEXT.other,cx=frame.x(xv),cy=frame.y(yv),r=active?6:4.5;if(active)frame.svg.appendChild(svgNode('circle',{cx,cy,r:r+3,fill:'none',stroke:'#c53f31','stroke-width':'2.5'}));let mark;if(style.shape==='circle')mark=svgNode('circle',{cx,cy,r});else{const pts=style.shape==='diamond'?[[cx,cy-r],[cx+r,cy],[cx,cy+r],[cx-r,cy]]:style.shape==='square'?[[cx-r,cy-r],[cx+r,cy-r],[cx+r,cy+r],[cx-r,cy+r]]:[[cx,cy-r],[cx+r,cy+r],[cx-r,cy+r]];mark=svgNode('polygon',{points:pts.map(p=>p.join(',')).join(' ')});}mark.setAttribute('fill',hollow?'#fff':style.color);mark.setAttribute('stroke',style.color);mark.setAttribute('stroke-width',hollow?'2':'1');mark.setAttribute('tabindex','0');mark.setAttribute('style','cursor:pointer');mark.appendChild(svgNode('title',{},`${s.sample} · ${style.label}: ${xv.toPrecision(4)}; ${yv.toPrecision(4)}`));mark.onclick=()=>activateChem(s.sample);frame.svg.appendChild(mark);if(active)frame.svg.appendChild(svgNode('text',{x:cx+9,y:cy-8,fill:'#76271f','font-size':'10','font-weight':'700'},s.sample));}
function renderTas(){const f=chartFrame(document.getElementById('tasChart'),{xDomain:[35,80],yDomain:[0,18],xTicks:[35,40,45,50,55,60,65,70,75,80],yTicks:[0,3,6,9,12,15,18],xLabel:'SiO₂, мас. %',yLabel:'Na₂O + K₂O, мас. %',height:500});f.logX=f.logY=false;
  TAS_FIELDS.forEach((field,i)=>{const poly=svgNode('polygon',{points:field.p.map(p=>`${f.x(p[0])},${f.y(p[1])}`).join(' '),fill:i%2?'#f5f7f5':'#edf2ee',stroke:'#98a49c','stroke-width':'.8'});f.svg.insertBefore(poly,f.svg.children[1]);const cx=field.p.reduce((a,p)=>a+p[0],0)/field.p.length,cy=field.p.reduce((a,p)=>a+p[1],0)/field.p.length;if(cx<=80)f.svg.appendChild(svgNode('text',{x:f.x(cx),y:f.y(cy)+3,'text-anchor':'middle',fill:'#68736c','font-size':'7.5'},field.label));});
  const selected=currentChemSamples();selected.filter(s=>s.sample!==chemActive).concat(selected.filter(s=>s.sample===chemActive)).forEach(s=>{const v=anhydrousMajor(s);plotPoint(f,s,v.SiO2,v.alkali);});
  const active=samples.find(s=>s.sample===chemActive&&s.geochem),note=document.getElementById('tasNote');if(active){const v=anhydrousMajor(active),field=TAS_FIELDS.find(a=>pointInPoly(v.SiO2,v.alkali,a.p));note.textContent=`${active.sample}: SiO₂ ${v.SiO2.toFixed(2)}%, Na₂O+K₂O ${v.alkali.toFixed(2)}%${field?` · поле «${field.label}»`:''}. ППП ${(active.geochem.major['ППП']||0).toFixed(2)}%.`;}}
function sampleElement(s,el){if(el==='K')return (s.geochem.major.K2O||0)*8301.5;if(el==='P')return (s.geochem.major.P2O5||0)*4364.2;if(el==='Ti')return (s.geochem.major.TiO2||0)*5993.4;return s.geochem.trace[el];}
function drawPattern(container,elements,normalizer){const selected=currentChemSamples(),palette=['#2c6da4','#b27718','#8a4f96','#278074','#a65242','#596a9a','#8d722f','#527b43','#7c5261','#3e7181','#765f9a','#95623b'],series=[];selected.forEach(s=>series.push({sample:s,values:elements.map(e=>sampleElement(s,e)/normalizer[e]),kind:'sample'}));CHEM_REFERENCES.filter(r=>chemRefSelected.has(r.name)).forEach((r,i)=>series.push({label:r.name,values:elements.map(e=>r.values[e]/normalizer[e]),kind:'reference',color:palette[i%palette.length]}));const vals=series.flatMap(a=>a.values).filter(v=>Number.isFinite(v)&&v>0),min=10**Math.floor(Math.log10(Math.min(...vals))),max=10**Math.ceil(Math.log10(Math.max(...vals))),ticks=[];for(let p=Math.log10(min);p<=Math.log10(max);p++)ticks.push(10**p);const f=chartFrame(container,{xDomain:[0,elements.length-1],yDomain:[min,max],xTicks:elements.map((e,i)=>({value:i,label:e})),yTicks:ticks,xLabel:'Элементы',yLabel:'Нормированное содержание',logY:true});f.logX=false;f.logY=true;
  const ordered=series.filter(a=>a.kind==='sample'&&a.sample.sample!==chemActive).concat(series.filter(a=>a.kind==='reference'),series.filter(a=>a.kind==='sample'&&a.sample.sample===chemActive));ordered.forEach(a=>{const active=a.sample?.sample===chemActive,context=a.sample?(CONTEXT[a.sample.geochem.context]||CONTEXT.other):null,color=active?'#c53f31':a.kind==='reference'?a.color:context.color,width=active?3:a.kind==='sample'?1.2:2,opacity=active||a.kind==='reference'?1:.38,dash=a.kind==='sample'?'':'6 4',points=a.values.map((v,i)=>Number.isFinite(v)&&v>0?`${f.x(i)},${f.y(v)}`:null);let d='';points.forEach((p,i)=>{if(p)d+=(i&&points[i-1]?'L':'M')+p;});const path=svgNode('path',{d,fill:'none',stroke:color,'stroke-width':width,opacity,'stroke-dasharray':dash,style:a.sample?'cursor:pointer':''});path.appendChild(svgNode('title',{},a.sample?.sample||a.label));if(a.sample)path.onclick=()=>activateChem(a.sample.sample);f.svg.appendChild(path);if(active)a.values.forEach((v,i)=>{if(Number.isFinite(v)&&v>0)f.svg.appendChild(svgNode('circle',{cx:f.x(i),cy:f.y(v),r:2.6,fill:color}));});});}
function basicFilter(s){const t=s.geochem.trace,ti=(s.geochem.major.TiO2||0)*5993.4;if(!(t.Nb>0&&t.Y>0&&t.Zr>0&&ti>0))return false;return .272*Math.log10(t.Nb/t.Y)-1.5528-Math.log10(t.Zr/ti)>=0;}
function renderLip(){const f=chartFrame(document.getElementById('lipChart'),{xDomain:[.1,10],yDomain:[.03,3],xTicks:[.1,.2,.5,1,2,5,10],yTicks:[.03,.1,.3,1,3],xLabel:'TiO₂/Yb',yLabel:'Th/Nb',logX:true,logY:true});f.logX=f.logY=true;const field=(x1,x2,y1,y2,fill,label,lx,ly)=>{f.svg.insertBefore(svgNode('polygon',{points:`${f.x(x1)},${f.y(y1)} ${f.x(x2)},${f.y(y1)} ${f.x(x2)},${f.y(y2)} ${f.x(x1)},${f.y(y2)}`,fill,stroke:'#6f7973','stroke-width':'1'}),f.svg.children[1]);f.svg.appendChild(svgNode('text',{x:f.x(lx),y:f.y(ly),'text-anchor':'middle',fill:'#59645d','font-size':'8'},label));};field(.2,.7,.15,3,'#f4e9ed','SZLM',.42,1.45);field(.2,5,.03,.15,'#e4f2f7','Plume array',2.8,.085);field(.7,5,.15,.2,'#fbf0df','EM-OIB',2,.173);f.svg.appendChild(svgNode('line',{x1:f.x(.7),y1:f.y(.03),x2:f.x(.7),y2:f.y(.15),stroke:'#6f7973'}));f.svg.appendChild(svgNode('text',{x:f.x(.36),y:f.y(.036),'text-anchor':'middle',fill:'#59645d','font-size':'7.5'},'MORB+OPB'));f.svg.appendChild(svgNode('text',{x:f.x(2.1),y:f.y(.036),'text-anchor':'middle',fill:'#59645d','font-size':'7.5'},'OIB+OPB'));const selected=currentChemSamples();selected.filter(s=>s.sample!==chemActive).concat(selected.filter(s=>s.sample===chemActive)).forEach(s=>{const t=s.geochem.trace,m=s.geochem.major;plotPoint(f,s,m.TiO2/t.Yb,t.Th/t.Nb,{hollow:!basicFilter(s)});});}
function renderTh(){const f=chartFrame(document.getElementById('thChart'),{xDomain:[.01,10],yDomain:[.01,100],xTicks:[.01,.1,1,10],yTicks:[.01,.1,1,10,100],xLabel:'Ta/Yb',yLabel:'Th/Yb',logX:true,logY:true});f.logX=f.logY=true;const xs=[.01,.03,.1,.3,1,3,10],band=svgNode('polygon',{points:xs.map(x=>`${f.x(x)},${f.y(2*x)}`).concat([...xs].reverse().map(x=>`${f.x(x)},${f.y(x)}`)).join(' '),fill:'#dce2de',opacity:'.75'});f.svg.insertBefore(band,f.svg.children[1]);f.svg.appendChild(svgNode('text',{x:f.x(.17),y:f.y(.25),fill:'#68736c','font-size':'8',transform:`rotate(-31 ${f.x(.17)} ${f.y(.25)})`},'мантийный массив'));const selected=currentChemSamples();selected.filter(s=>s.sample!==chemActive).concat(selected.filter(s=>s.sample===chemActive)).forEach(s=>{const t=s.geochem.trace;plotPoint(f,s,t.Ta/t.Yb,t.Th/t.Yb);});}
function renderChemSidebar(){const available=samples.filter(s=>s.geochem),box=document.getElementById('chemSamples');document.getElementById('chemCount').textContent=`Выбрано ${chemSelected.size} из ${available.length}`;document.getElementById('chemActiveTitle').textContent=chemActive?`Активный: ${chemActive}`:'Выбранный образец';box.innerHTML=available.map(s=>`<button class="chem-sample ${chemSelected.has(s.sample)?'selected':''} ${s.sample===chemActive?'active':''}" data-sample="${esc(s.sample)}"><i class="chem-check"></i><b>${esc(s.sample)}</b><span class="chem-group">${esc((CONTEXT[s.geochem.context]||CONTEXT.other).label)}</span></button>`).join('');box.querySelectorAll('button').forEach(b=>b.onclick=()=>{const id=b.dataset.sample;if(chemSelected.has(id)){if(chemActive!==id)chemActive=id;else if(chemSelected.size>1){chemSelected.delete(id);chemActive=[...chemSelected][0];}}else{chemSelected.add(id);chemActive=id;}renderChemSidebar();renderChemCharts();});}
function renderReferencePicker(){const box=document.getElementById('chemRefs');box.innerHTML=CHEM_REFERENCES.map((r,i)=>`<label class="chem-ref"><input type="checkbox" data-reference="${esc(r.name)}" ${chemRefSelected.has(r.name)?'checked':''}><span>${esc(r.name)}</span></label>`).join('');box.querySelectorAll('input').forEach(input=>input.onchange=()=>{if(input.checked)chemRefSelected.add(input.dataset.reference);else chemRefSelected.delete(input.dataset.reference);renderChemCharts();});}
function renderChemCharts(){if(!chemSelected.size)return;renderTas();drawPattern(document.getElementById('reeChart'),REE_ELEMENTS,C1);drawPattern(document.getElementById('spiderChart'),SPIDER_ELEMENTS,PM);renderLip();renderTh();}
function openChemistry(sampleId){const available=samples.filter(s=>s.geochem);if(!available.length)return;const chosen=available.find(s=>s.sample===sampleId)||available.find(s=>s.sample===activeSample)||available[0];chemActive=chosen.sample;chemSelected=new Set([chosen.sample]);if(!chemRefSelected.size){const defaults=CHEM_REFERENCES.filter(r=>r.name.startsWith('N-MORB (Gale')||r.name.startsWith('OIB (Sun')).map(r=>r.name);chemRefSelected=new Set(defaults);}const modal=document.getElementById('chemModal');modal.classList.add('visible');modal.setAttribute('aria-hidden','false');renderChemSidebar();renderReferencePicker();requestAnimationFrame(renderChemCharts);}
function closeChemistry(){const modal=document.getElementById('chemModal');modal.classList.remove('visible');modal.setAttribute('aria-hidden','true');}
const UPB_L238=1.55125e-10,UPB_L235=9.8485e-10,UPB_H=30;
let upbSample=null,upbView=null,upbDrag=null,upbPinch=null,upbComparison=false;
const upbPointers=new Map();
function upbDiscordance(a,method){if(method==='d1')return 100*(a.age_75/a.age_68-1);if(method==='d2')return 100*(a.age_76/a.age_68-1);return a.age_68<1000?100*(a.age_75/a.age_68-1):100*(a.age_76/a.age_68-1);}
function upbRowsFrom(analyses){const method=document.getElementById('upbMethod').value,limit=Number(document.getElementById('upbLimit').value),context=document.getElementById('upbContext').checked;return (analyses||[]).map(a=>{const disc=upbDiscordance(a,method),bestAge=a.age_68<1000?a.age_68:a.age_76,bestErr=a.age_68<1000?a.err_age_68:a.err_age_76,analytical=Math.abs(disc)<=limit;return{...a,disc,bestAge,bestErr,analytical,accepted:analytical&&(!context||!a.context_excluded)};});}
function upbRows(){return upbSample?.upb?upbRowsFrom(upbSample.upb.analyses):[];}
function upbComparisonSeries(){const sm=samples.find(s=>s.sample==='SM25-07')?.upb,lit=GEOCHRONOLOGY.datasets.find(d=>d.id==='KOCHNEV-1812-1-SILL');return[
  {id:'SM25-07',label:'SM25-07',color:'#126e64',rows:upbRowsFrom(sm?.analyses)},
  {id:'1812-1',label:'1812-1 · Кочнев',color:'#6b5a92',rows:upbRowsFrom(lit?.upb_analyses)}
];}
function upbComparisonRows(){return upbComparisonSeries().flatMap(s=>s.rows.map(a=>({...a,seriesId:s.id,seriesLabel:s.label,seriesColor:s.color})));}
function upbActiveRows(){return upbComparison?upbComparisonRows():upbRows();}
function upbFormula(){const m=document.getElementById('upbMethod').value;if(m==='d1')return 'D1 = 100 × (t²⁰⁷Pb/²³⁵U ÷ t²⁰⁶Pb/²³⁸U − 1)';if(m==='d2')return 'D2 = 100 × (t²⁰⁷Pb/²⁰⁶Pb ÷ t²⁰⁶Pb/²³⁸U − 1)';return 'D1 при t²⁰⁶Pb/²³⁸U < 1000 млн лет; D2 при ≥ 1000 млн лет';}
function linTicks(lo,hi,n=6){const raw=(hi-lo)/Math.max(1,n-1),pow=10**Math.floor(Math.log10(raw)),q=raw/pow,step=(q<=1?1:q<=2?2:q<=5?5:10)*pow,start=Math.ceil(lo/step)*step,out=[];for(let v=start;v<=hi+step*.001;v+=step)out.push(v);return out;}
function tickText(v,span){if(Math.abs(v)>=100)return v.toFixed(0);if(span<.05)return v.toFixed(3);if(span<1)return v.toFixed(2);return v.toFixed(1);}
function upbSvgFrame(container,domain,height=380){container.innerHTML='';const width=Math.max(480,container.clientWidth||900),m={l:68,r:18,t:14,b:52},pw=width-m.l-m.r,ph=height-m.t-m.b,svg=svgNode('svg',{viewBox:`0 0 ${width} ${height}`,role:'img'}),x=v=>m.l+(v-domain.x0)/(domain.x1-domain.x0)*pw,y=v=>m.t+ph-(v-domain.y0)/(domain.y1-domain.y0)*ph;container.appendChild(svg);const grid=svgNode('g');svg.appendChild(grid);linTicks(domain.y0,domain.y1).forEach(t=>{const py=y(t);grid.appendChild(svgNode('line',{x1:m.l,y1:py,x2:m.l+pw,y2:py,stroke:'#dfe4e0'}));grid.appendChild(svgNode('text',{x:m.l-8,y:py+3,'text-anchor':'end',fill:'#68736c','font-size':'9'},tickText(t,domain.y1-domain.y0)));});linTicks(domain.x0,domain.x1).forEach(t=>{const px=x(t);grid.appendChild(svgNode('line',{x1:px,y1:m.t,x2:px,y2:m.t+ph,stroke:'#edf0ee'}));grid.appendChild(svgNode('text',{x:px,y:m.t+ph+17,'text-anchor':'middle',fill:'#68736c','font-size':'9'},tickText(t,domain.x1-domain.x0)));});svg.appendChild(svgNode('line',{x1:m.l,y1:m.t+ph,x2:m.l+pw,y2:m.t+ph,stroke:'#59645d'}));svg.appendChild(svgNode('line',{x1:m.l,y1:m.t,x2:m.l,y2:m.t+ph,stroke:'#59645d'}));return{svg,width,height,m,pw,ph,x,y,domain};}
function upbWholeView(){return{x0:0,x1:32,y0:0,y1:.75};}
function upbFitView(rows){const pts=rows.filter(a=>a.accepted);if(!pts.length)return upbWholeView();let x0=Math.min(...pts.map(a=>a.ratio_75-2*a.err_ratio_75)),x1=Math.max(...pts.map(a=>a.ratio_75+2*a.err_ratio_75)),y0=Math.min(...pts.map(a=>a.ratio_68-2*a.err_ratio_68)),y1=Math.max(...pts.map(a=>a.ratio_68+2*a.err_ratio_68)),xp=Math.max(.12,(x1-x0)*.18),yp=Math.max(.008,(y1-y0)*.18);return{x0:Math.max(0,x0-xp),x1:x1+xp,y0:Math.max(0,y0-yp),y1:y1+yp};}
function upbZoomedView(base,factor,cx=.5,cy=.5){const wx=base.x1-base.x0,wy=base.y1-base.y0,x=base.x0+cx*wx,y=base.y1-cy*wy,nx=Math.max(.02,Math.min(40,wx*factor)),ny=Math.max(.001,Math.min(1,wy*factor));return{x0:x-cx*nx,x1:x+(1-cx)*nx,y0:y-(1-cy)*ny,y1:y+cy*ny};}
function upbScaleView(factor,cx=.5,cy=.5){upbView=upbZoomedView(upbView||upbWholeView(),factor,cx,cy);renderUpbConcordia();}
function renderUpbComparisonConcordia(){const rows=upbComparisonRows(),container=document.getElementById('upbConcordia'),domain=upbView||upbWholeView(),height=Math.max(430,Math.min(650,(container.clientWidth||900)*.70)),f=upbSvgFrame(container,domain,height);f.svg.setAttribute('aria-label','Сравнение SM25-07 и 1812-1 на Wetherill concordia');const clipId='upbCompareClip',defs=svgNode('defs'),clip=svgNode('clipPath',{id:clipId});clip.appendChild(svgNode('rect',{x:f.m.l,y:f.m.t,width:f.pw,height:f.ph}));defs.appendChild(clip);f.svg.appendChild(defs);const plot=svgNode('g',{'clip-path':`url(#${clipId})`});f.svg.appendChild(plot);let d='';for(let age=1;age<=3500;age+=3){const r68=Math.expm1(UPB_L238*age*1e6),r75=Math.expm1(UPB_L235*age*1e6);d+=(d?'L':'M')+f.x(r75).toFixed(2)+','+f.y(r68).toFixed(2);}plot.appendChild(svgNode('path',{d,fill:'none',stroke:'#26322d','stroke-width':'2'}));const ageStep=(domain.x1-domain.x0)<3?25:(domain.x1-domain.x0)<10?100:200;for(let age=ageStep;age<=3400;age+=ageStep){const r68=Math.expm1(UPB_L238*age*1e6),r75=Math.expm1(UPB_L235*age*1e6);if(r75<domain.x0||r75>domain.x1||r68<domain.y0||r68>domain.y1)continue;plot.appendChild(svgNode('circle',{cx:f.x(r75),cy:f.y(r68),r:2.2,fill:'#26322d'}));plot.appendChild(svgNode('text',{x:f.x(r75)+4,y:f.y(r68)-4,fill:'#53605a','font-size':'8'},String(age)));}const sx=f.pw/(domain.x1-domain.x0),sy=-f.ph/(domain.y1-domain.y0);rows.filter(a=>!a.accepted).concat(rows.filter(a=>a.accepted)).forEach(a=>{const color=a.accepted?a.seriesColor:'#aeb4b0',opacity=a.accepted?.34:.16,sxx=(a.err_ratio_75*sx)**2,syy=(a.err_ratio_68*sy)**2,sxy=(a.rho_3*a.err_ratio_75*a.err_ratio_68)*sx*sy,tr=sxx+syy,det=sxx*syy-sxy*sxy,root=Math.sqrt(Math.max(0,tr*tr/4-det)),l1=tr/2+root,l2=tr/2-root,angle=.5*Math.atan2(2*sxy,sxx-syy)*180/Math.PI,e=svgNode('ellipse',{cx:f.x(a.ratio_75),cy:f.y(a.ratio_68),rx:2*Math.sqrt(Math.max(0,l1)),ry:2*Math.sqrt(Math.max(0,l2)),transform:`rotate(${angle} ${f.x(a.ratio_75)} ${f.y(a.ratio_68)})`,fill:color,stroke:color,'stroke-width':a.accepted?'1.2':'.7',opacity});e.appendChild(svgNode('title',{},`${a.seriesLabel} · ${a.sid}: D ${a.disc.toFixed(1)}%`));plot.appendChild(e);});const legendX=f.m.l+12,legendY=f.m.t+18;upbComparisonSeries().forEach((s,i)=>{f.svg.appendChild(svgNode('line',{x1:legendX,y1:legendY+i*20,x2:legendX+22,y2:legendY+i*20,stroke:s.color,'stroke-width':'5'}));f.svg.appendChild(svgNode('text',{x:legendX+30,y:legendY+3+i*20,fill:'#465149','font-size':'10'},`${s.label} · ${s.rows.filter(a=>a.accepted).length}/${s.rows.length}`));});f.svg.appendChild(svgNode('text',{x:f.m.l+f.pw/2,y:f.height-8,'text-anchor':'middle',fill:'#465149','font-size':'10'},'²⁰⁷Pb/²³⁵U'));f.svg.appendChild(svgNode('text',{x:13,y:f.m.t+f.ph/2,'text-anchor':'middle',fill:'#465149','font-size':'10',transform:`rotate(-90 13 ${f.m.t+f.ph/2})`},'²⁰⁶Pb/²³⁸U'));}
function renderUpbConcordia(){if(upbComparison)return renderUpbComparisonConcordia();const rows=upbRows(),container=document.getElementById('upbConcordia'),domain=upbView||upbWholeView(),height=Math.max(420,Math.min(620,(container.clientWidth||900)*.68)),f=upbSvgFrame(container,domain,height);f.svg.setAttribute('aria-label','Wetherill concordia: X 207Pb/235U, Y 206Pb/238U');const clipId='upbPlotClip',defs=svgNode('defs'),clip=svgNode('clipPath',{id:clipId});clip.appendChild(svgNode('rect',{x:f.m.l,y:f.m.t,width:f.pw,height:f.ph}));defs.appendChild(clip);f.svg.appendChild(defs);const plot=svgNode('g',{'clip-path':`url(#${clipId})`});f.svg.appendChild(plot);let d='';for(let age=1;age<=3500;age+=4){const r68=Math.expm1(UPB_L238*age*1e6),r75=Math.expm1(UPB_L235*age*1e6),px=f.x(r75),py=f.y(r68);d+=(d?'L':'M')+px.toFixed(2)+','+py.toFixed(2);}plot.appendChild(svgNode('path',{d,fill:'none',stroke:'#26322d','stroke-width':'2'}));const ageStep=(domain.x1-domain.x0)<3?50:(domain.x1-domain.x0)<10?100:200;for(let age=ageStep;age<=3400;age+=ageStep){const r68=Math.expm1(UPB_L238*age*1e6),r75=Math.expm1(UPB_L235*age*1e6);if(r75<domain.x0||r75>domain.x1||r68<domain.y0||r68>domain.y1)continue;plot.appendChild(svgNode('circle',{cx:f.x(r75),cy:f.y(r68),r:2.3,fill:'#26322d'}));plot.appendChild(svgNode('text',{x:f.x(r75)+4,y:f.y(r68)-4,fill:'#53605a','font-size':'8'},String(age)));}const sx=f.pw/(domain.x1-domain.x0),sy=-f.ph/(domain.y1-domain.y0);rows.filter(a=>!a.accepted).concat(rows.filter(a=>a.accepted)).forEach(a=>{const color=a.accepted?'#167d70':a.analytical?'#d77b28':'#aeb4b0',opacity=a.accepted?.42:.28,sxx=(a.err_ratio_75*sx)**2,syy=(a.err_ratio_68*sy)**2,sxy=(a.rho_3*a.err_ratio_75*a.err_ratio_68)*sx*sy,tr=sxx+syy,det=sxx*syy-sxy*sxy,root=Math.sqrt(Math.max(0,tr*tr/4-det)),l1=tr/2+root,l2=tr/2-root,angle=.5*Math.atan2(2*sxy,sxx-syy)*180/Math.PI,e=svgNode('ellipse',{cx:f.x(a.ratio_75),cy:f.y(a.ratio_68),rx:2*Math.sqrt(Math.max(0,l1)),ry:2*Math.sqrt(Math.max(0,l2)),transform:`rotate(${angle} ${f.x(a.ratio_75)} ${f.y(a.ratio_68)})`,fill:color,stroke:color,'stroke-width':a.accepted?'1':'.7',opacity});e.appendChild(svgNode('title',{},`${a.sid}: ²⁰⁷Pb/²³⁵U ${a.ratio_75.toFixed(4)}; ²⁰⁶Pb/²³⁸U ${a.ratio_68.toFixed(5)}; D ${a.disc.toFixed(1)}%`));plot.appendChild(e);});f.svg.appendChild(svgNode('text',{x:f.m.l+f.pw/2,y:f.height-8,'text-anchor':'middle',fill:'#465149','font-size':'10'},'²⁰⁷Pb/²³⁵U'));f.svg.appendChild(svgNode('text',{x:13,y:f.m.t+f.ph/2,'text-anchor':'middle',fill:'#465149','font-size':'10',transform:`rotate(-90 13 ${f.m.t+f.ph/2})`},'²⁰⁶Pb/²³⁸U'));}
function upbKdeValues(ages,xs){const c=1/(Math.max(1,ages.length)*UPB_H*Math.sqrt(2*Math.PI));return xs.map(x=>c*ages.reduce((s,a)=>s+Math.exp(-.5*((x-a)/UPB_H)**2),0));}
function renderUpbKde(){const rows=upbRows(),accepted=rows.filter(a=>a.accepted),ages=accepted.map(a=>a.bestAge),container=document.getElementById('upbKde');if(!ages.length){container.textContent='Нет принятых анализов';return;}const amin=Math.min(...ages),amax=Math.max(...ages),pad=Math.max(100,(amax-amin)*.08),lo=Math.max(0,Math.floor((amin-pad)/50)*50),hi=Math.ceil((amax+pad)/50)*50,steps=Math.max(400,Math.round((hi-lo)*2)),xs=Array.from({length:steps+1},(_,i)=>lo+(hi-lo)*i/steps),dens=upbKdeValues(ages,xs),peaks=[];for(let i=1;i<dens.length-1;i++)if(dens[i]>dens[i-1]&&dens[i]>=dens[i+1]){const y0=dens[i-1],y1=dens[i],y2=dens[i+1],den=y0-2*y1+y2,offset=den?Math.max(-.5,Math.min(.5,.5*(y0-y2)/den)):0;peaks.push({age:xs[i]+offset*(xs[i+1]-xs[i]),d:y1});}const bins=[];for(let x=lo;x<hi;x+=50)bins.push({x,n:ages.filter(a=>a>=x&&a<x+50).length});const maxD=Math.max(...dens)*1.18,maxN=Math.max(1,...bins.map(b=>b.n)),f=upbSvgFrame(container,{x0:lo,x1:hi,y0:0,y1:maxD},390);f.svg.setAttribute('aria-label',`KDE при h = ${UPB_H} млн лет; рассчитано пиков: ${peaks.length}`);bins.forEach(b=>{const x=f.x(b.x)+1,w=Math.max(1,f.x(b.x+50)-f.x(b.x)-2),h=b.n/maxN*f.ph*.42;f.svg.appendChild(svgNode('rect',{x,y:f.m.t+f.ph-h,width:w,height:h,fill:'#7fa7bd',opacity:'.28'}));});let line='';xs.forEach((x,i)=>{line+=(i?'L':'M')+f.x(x)+','+f.y(dens[i]);});f.svg.appendChild(svgNode('path',{d:line,fill:'none',stroke:'#126e64','stroke-width':'2.7'}));rows.forEach(a=>{const x=f.x(a.bestAge),color=a.accepted?'#26322d':a.analytical?'#d77b28':'#b05a55';f.svg.appendChild(svgNode('line',{x1:x,y1:f.m.t+f.ph+(a.accepted?3:14),x2:x,y2:f.m.t+f.ph+(a.accepted?12:24),stroke:color,'stroke-width':a.accepted?'1':'1.5',opacity:'.72'}));});peaks.forEach((p,i)=>{f.svg.appendChild(svgNode('circle',{cx:f.x(p.age),cy:f.y(p.d),r:3,fill:'#126e64'}));f.svg.appendChild(svgNode('text',{x:f.x(p.age),y:f.y(p.d)-(i%2?15:7),'text-anchor':'middle',fill:'#465149','font-size':'8'},p.age.toFixed(0)));});f.svg.appendChild(svgNode('text',{x:f.m.l+f.pw/2,y:f.height-8,'text-anchor':'middle',fill:'#465149','font-size':'10'},'Best age, млн лет'));document.getElementById('upbPeaks').innerHTML=`<b>Все локальные пики KDE (${peaks.length}, h = ${UPB_H}):</b> `+peaks.map(p=>`<span>${p.age.toFixed(0)} млн лет</span>`).join('');}
function upbWeighted(rows){const accepted=rows.filter(a=>a.accepted),weights=accepted.map(a=>1/a.bestErr**2),sw=weights.reduce((a,b)=>a+b,0),mean=accepted.reduce((s,a,i)=>s+weights[i]*a.bestAge,0)/sw,u2=2/Math.sqrt(sw),mswd=accepted.length>1?accepted.reduce((s,a,i)=>s+weights[i]*(a.bestAge-mean)**2,0)/(accepted.length-1):NaN;return{accepted,mean,u2,mswd};}
function renderUpbComparisonWm(){const container=document.getElementById('upbWm'),note=document.getElementById('upbWmNote'),states=upbComparisonSeries().map(s=>({...s,...upbWeighted(s.rows)})),width=Math.max(480,container.clientWidth||620),height=310,m={l:125,r:28,t:35,b:58},pw=width-m.l-m.r,values=states.flatMap(s=>[s.mean-s.u2,s.mean+s.u2]),lo=Math.floor((Math.min(...values)-8)/10)*10,hi=Math.ceil((Math.max(...values)+8)/10)*10,x=v=>m.l+(v-lo)/(hi-lo)*pw,svg=svgNode('svg',{viewBox:`0 0 ${width} ${height}`,role:'img','aria-label':'Сравнение weighted mean для SM25-07 и 1812-1'});container.innerHTML='';container.appendChild(svg);for(let t=Math.ceil(lo/10)*10;t<=hi;t+=10){svg.appendChild(svgNode('line',{x1:x(t),y1:m.t,x2:x(t),y2:height-m.b,stroke:'#dfe4e0'}));svg.appendChild(svgNode('text',{x:x(t),y:height-m.b+22,'text-anchor':'middle',fill:'#68736c','font-size':'10'},String(t)));}states.forEach((s,i)=>{const y=m.t+58+i*88;svg.appendChild(svgNode('text',{x:m.l-12,y:y+4,'text-anchor':'end',fill:'#465149','font-size':'11'},s.label));svg.appendChild(svgNode('line',{x1:x(s.mean-s.u2),y1:y,x2:x(s.mean+s.u2),y2:y,stroke:s.color,'stroke-width':'4'}));svg.appendChild(svgNode('line',{x1:x(s.mean-s.u2),y1:y-8,x2:x(s.mean-s.u2),y2:y+8,stroke:s.color,'stroke-width':'2'}));svg.appendChild(svgNode('line',{x1:x(s.mean+s.u2),y1:y-8,x2:x(s.mean+s.u2),y2:y+8,stroke:s.color,'stroke-width':'2'}));svg.appendChild(svgNode('circle',{cx:x(s.mean),cy:y,r:6,fill:s.color,stroke:'white','stroke-width':'1.5'}));svg.appendChild(svgNode('text',{x:x(s.mean),y:y-15,'text-anchor':'middle',fill:'#465149','font-size':'10'},`${s.mean.toFixed(1)} ± ${s.u2.toFixed(1)} млн лет`));svg.appendChild(svgNode('text',{x:x(s.mean),y:y+22,'text-anchor':'middle',fill:'#68736c','font-size':'9'},`MSWD ${s.mswd.toFixed(2)} · n = ${s.accepted.length}`));});svg.appendChild(svgNode('line',{x1:m.l,y1:height-m.b,x2:m.l+pw,y2:height-m.b,stroke:'#59645d'}));svg.appendChild(svgNode('text',{x:m.l+pw/2,y:height-10,'text-anchor':'middle',fill:'#465149','font-size':'10'},'²⁰⁶Pb/²³⁸U weighted mean, млн лет (2σ внутр.)'));const a=states[0],b=states[1],delta=b.mean-a.mean,delta2=Math.hypot(a.u2,b.u2);note.textContent=`Разница 1812-1 − SM25-07: ${delta.toFixed(1)} ± ${delta2.toFixed(1)} млн лет (2σ, внутренние ошибки).`;}
function renderUpbWm(){if(upbComparison)return renderUpbComparisonWm();const rows=upbRows(),accepted=rows.filter(a=>a.accepted).sort((a,b)=>a.bestAge-b.bestAge),container=document.getElementById('upbWm'),note=document.getElementById('upbWmNote');if(upbSample.sample==='SM25-03'){container.innerHTML='<div class="photo-empty">Единая магматическая популяция не выделена. Weighted mean для всей смешанной выборки геологически некорректно.</div>';note.textContent='Сначала нужно обосновать выбор кластера по CL/BSE-контексту спотов.';return;}if(accepted.length<2){container.textContent='Недостаточно принятых анализов';note.textContent='';return;}const weights=accepted.map(a=>1/a.bestErr**2),sw=weights.reduce((a,b)=>a+b,0),mean=accepted.reduce((s,a,i)=>s+weights[i]*a.bestAge,0)/sw,u2=2/Math.sqrt(sw),mswd=accepted.reduce((s,a,i)=>s+weights[i]*(a.bestAge-mean)**2,0)/(accepted.length-1),x0=Math.min(...accepted.map(a=>a.bestAge-2*a.bestErr),mean-u2),x1=Math.max(...accepted.map(a=>a.bestAge+2*a.bestErr),mean+u2),pad=Math.max(5,(x1-x0)*.12),height=Math.max(350,Math.min(900,accepted.length*10+70)),f=upbSvgFrame(container,{x0:x0-pad,x1:x1+pad,y0:.5,y1:accepted.length+.5},height);f.svg.setAttribute('aria-label','Weighted mean: возраст по горизонтальной оси, анализы по вертикальной');f.svg.insertBefore(svgNode('rect',{x:f.x(mean-u2),y:f.m.t,width:f.x(mean+u2)-f.x(mean-u2),height:f.ph,fill:'#a9d5cf',opacity:'.62'}),f.svg.children[1]);f.svg.appendChild(svgNode('line',{x1:f.x(mean),y1:f.m.t,x2:f.x(mean),y2:f.m.t+f.ph,stroke:'#bb3d3d','stroke-width':'1.6','stroke-dasharray':'6 4'}));accepted.forEach((a,i)=>{const y=f.y(i+1),lo=f.x(a.bestAge-2*a.bestErr),hi=f.x(a.bestAge+2*a.bestErr);f.svg.appendChild(svgNode('line',{x1:lo,y1:y,x2:hi,y2:y,stroke:'#438d85'}));const c=svgNode('circle',{cx:f.x(a.bestAge),cy:y,r:3.3,fill:'#126e64'});c.appendChild(svgNode('title',{},`${a.sid}: ${a.bestAge.toFixed(1)} ± ${(2*a.bestErr).toFixed(1)} млн лет (2σ)`));f.svg.appendChild(c);});f.svg.appendChild(svgNode('text',{x:f.m.l+f.pw/2,y:f.height-8,'text-anchor':'middle',fill:'#465149','font-size':'10'},'Best age, млн лет'));const kind=upbSample.upb.sample_kind;let caveat='';if(kind==='detrital'||upbSample.sample==='SM25-16')caveat=' · Для детритового спектра это описательная величина, не возраст образца.';note.textContent=`${mean.toFixed(1)} ± ${u2.toFixed(1)} млн лет (2σ, внутр.); MSWD = ${mswd.toFixed(2)}; n = ${accepted.length}`+caveat;}
function renderUpb(){if(!upbSample&&!upbComparison)return;const rows=upbComparison?upbComparisonRows():upbRows(),accepted=rows.filter(a=>a.accepted),analyticalRejected=rows.filter(a=>!a.analytical),contextRejected=rows.filter(a=>a.analytical&&!a.accepted),limit=Number(document.getElementById('upbLimit').value),magmatic=upbComparison||['SM25-07','1812-1'].includes(upbSample?.sample);document.getElementById('upbGrid').classList.toggle('magmatic',magmatic);document.getElementById('upbCompare').style.display=!upbComparison&&['SM25-07','1812-1'].includes(upbSample?.sample)?'block':'none';document.getElementById('upbLimitLabel').textContent=`${limit}%`;document.querySelectorAll('.upb-preset').forEach(b=>b.classList.toggle('active',Number(b.dataset.limit)===limit));if(upbComparison){const detail=upbComparisonSeries().map(s=>`${s.label}: ${s.rows.filter(a=>a.accepted).length}/${s.rows.length}`).join(' · ');document.getElementById('upbStatus').textContent=`${upbFormula()} · ${detail}. Для 1812-1 два определения вне кластера скрыты как контекстные исключения.`;}else document.getElementById('upbStatus').textContent=`${upbFormula()} · принято ${accepted.length} из ${rows.length}; по дискордантности исключено ${analyticalRejected.length}`+(contextRejected.length?`; контекстно исключено ${contextRejected.length}`:'')+'. Методы Вирмиша и Павермана добавим после сверки формул по статьям.';renderUpbConcordia();if(!magmatic)renderUpbKde();renderUpbWm();}
function showUpbModal(){const modal=document.getElementById('upbModal');modal.classList.add('visible');modal.setAttribute('aria-hidden','false');requestAnimationFrame(renderUpb);}
function openUpb(sampleId){const s=samples.find(x=>x.sample===sampleId&&x.upb);if(!s)return;upbComparison=false;upbSample=s;upbView=upbFitView(upbRows());document.getElementById('upbTitle').textContent=`${s.sample} · U–Pb`;showUpbModal();}
function openLiteratureUpb(datasetId){const d=GEOCHRONOLOGY.datasets.find(x=>x.id===datasetId&&x.upb_analyses);if(!d)return;upbComparison=false;upbSample={sample:d.label,upb:{analyses:d.upb_analyses,sample_kind:'igneous'}};document.getElementById('upbContext').checked=true;upbView=upbFitView(upbRows());document.getElementById('upbTitle').textContent=`${d.label} · U–Pb`;showUpbModal();}
function openUpbComparison(){upbComparison=true;upbSample=null;document.getElementById('upbContext').checked=true;upbView=upbFitView(upbComparisonRows());document.getElementById('upbTitle').textContent='1812-1 и SM25-07 · сравнение U–Pb';showUpbModal();}
function closeUpb(){document.getElementById('upbModal').classList.remove('visible');document.getElementById('upbModal').setAttribute('aria-hidden','true');upbComparison=false;upbDrag=null;upbPinch=null;upbPointers.clear();}
let geoFocus=null;
const geoSelections=new Map();
const geoVisibility=new Map();
function geoKdeValues(ages,xs,h){const c=1/(Math.max(1,ages.length)*h*Math.sqrt(2*Math.PI));return xs.map(x=>c*ages.reduce((sum,age)=>sum+Math.exp(-.5*((x-age)/h)**2),0));}
function geoFeatures(xs,dens){const maxima=[],minima=[];for(let i=1;i<dens.length-1;i++){if(dens[i]>dens[i-1]&&dens[i]>=dens[i+1])maxima.push(i);if(dens[i]<=dens[i-1]&&dens[i]<dens[i+1])minima.push(i);}return maxima.map(i=>{const left=[...minima].reverse().find(j=>j<i),right=minima.find(j=>j>i),y0=dens[i-1],y1=dens[i],y2=dens[i+1],den=y0-2*y1+y2,offset=den?Math.max(-.5,Math.min(.5,.5*(y0-y2)/den)):0;return{i,age:xs[i]+offset*(xs[i+1]-xs[i]),d:y1,lo:left===undefined?xs[0]:xs[left],hi:right===undefined?xs[xs.length-1]:xs[right]};});}
function geoLogGamma(z){const p=[.9999999999998099,676.5203681218851,-1259.1392167224028,771.3234287776531,-176.6150291621406,12.5073432786869,-.13857109526572,9.984369578019572e-6,1.5056327351493116e-7];if(z<.5)return Math.log(Math.PI)-Math.log(Math.sin(Math.PI*z))-geoLogGamma(1-z);z-=1;let x=p[0];for(let i=1;i<p.length;i++)x+=p[i]/(z+i);const t=z+7.5;return .5*Math.log(2*Math.PI)+(z+.5)*Math.log(t)-t+Math.log(x);}
function geoGammaQ(a,x){if(x<0||a<=0)return NaN;if(x===0)return 1;const gln=geoLogGamma(a),eps=3e-14,tiny=1e-300,maxIt=200;if(x<a+1){let ap=a,sum=1/a,del=sum;for(let n=1;n<=maxIt;n++){ap+=1;del*=x/ap;sum+=del;if(Math.abs(del)<Math.abs(sum)*eps)break;}return Math.max(0,Math.min(1,1-sum*Math.exp(-x+a*Math.log(x)-gln)));}let b=x+1-a,c=1/tiny,d=1/b,h=d;for(let i=1;i<=maxIt;i++){const an=-i*(i-a);b+=2;d=an*d+b;if(Math.abs(d)<tiny)d=tiny;c=b+an/c;if(Math.abs(c)<tiny)c=tiny;d=1/d;const del=d*c;h*=del;if(Math.abs(del-1)<eps)break;}return Math.max(0,Math.min(1,Math.exp(-x+a*Math.log(x)-gln)*h));}
function geoWeighted(rows){if(!rows.length)return null;const weights=rows.map(a=>1/(a.err_1s*a.err_1s)),sw=weights.reduce((a,b)=>a+b,0),mean=rows.reduce((sum,a,i)=>sum+weights[i]*a.age,0)/sw,u2=2/Math.sqrt(sw);if(rows.length===1)return{mean,u2,expanded:u2,mswd:NaN,p:NaN,n:1};const chi2=rows.reduce((sum,a,i)=>sum+weights[i]*(a.age-mean)**2,0),df=rows.length-1,mswd=chi2/df;return{mean,u2,expanded:u2*Math.sqrt(Math.max(1,mswd)),mswd,p:geoGammaQ(df/2,chi2/2),n:rows.length};}
function geoSeries(dataset){const base={id:dataset.id,label:dataset.source,source:dataset.source,color:'#126e64',dash:false,order:0,analyses:dataset.analyses};return dataset.comparison?[base,{...dataset.comparison,color:'#6b5a92',dash:true,order:1}]:[base];}
function geoDomain(){const all=GEOCHRONOLOGY.datasets.flatMap(d=>geoSeries(d).flatMap(s=>s.analyses.filter(a=>a.accepted).map(a=>a.age))),hi=Math.max(3500,Math.ceil(Math.max(...all,3300)/500)*500);return{x0:0,x1:hi};}
function geoYoungDomain(dataset){return dataset.id==='GUBINSKAYA-COMPARISON'?{x0:650,x1:950}:{x0:500,x1:1200};}
function geoVisibleSeries(dataset){const all=geoSeries(dataset);if(!geoVisibility.has(dataset.id))geoVisibility.set(dataset.id,new Set(all.map(s=>s.id)));const visible=geoVisibility.get(dataset.id);return all.filter(s=>visible.has(s.id));}
function geoStates(seriesList,domain,h){const steps=Math.max(400,Math.round((domain.x1-domain.x0)*2)),xs=Array.from({length:steps+1},(_,i)=>domain.x0+(domain.x1-domain.x0)*i/steps);return seriesList.map(series=>{const accepted=series.analyses.filter(a=>a.accepted&&Number.isFinite(a.age)&&Number.isFinite(a.err_1s)&&a.err_1s>0),dens=geoKdeValues(accepted.map(a=>a.age),xs,h);return{series,accepted,xs,dens,peaks:geoFeatures(xs,dens)};});}
function geoSelected(states,previous){if(!previous)return null;const state=states.find(s=>s.series.id===previous.seriesId);if(!state?.peaks.length)return null;const index=state.peaks.reduce((best,p,i)=>Math.abs(p.age-previous.age)<Math.abs(state.peaks[best].age-previous.age)?i:best,0);return{state,index};}
function geoChart(dataset,container,states,selectedState,domain,h,height){const maxD=Math.max(...states.flatMap(s=>s.dens),1e-8)*1.14,f=upbSvgFrame(container,{...domain,y0:0,y1:maxD},height);f.svg.setAttribute('aria-label',`${dataset.label}; объект: ${dataset.object}; KDE ${domain.x0}–${domain.x1} млн лет, ширина ядра ${h} млн лет`);const chosen=selectedState?selectedState.state.peaks[selectedState.index]:null;if(chosen){f.svg.insertBefore(svgNode('rect',{x:f.x(chosen.lo),y:f.m.t,width:Math.max(1,f.x(chosen.hi)-f.x(chosen.lo)),height:f.ph,fill:'#b9d8cb',opacity:'.42'}),f.svg.children[1]);}states.forEach(state=>{let path='';state.xs.forEach((x,i)=>{path+=(i?'L':'M')+f.x(x).toFixed(2)+','+f.y(state.dens[i]).toFixed(2);});f.svg.appendChild(svgNode('path',{d:path,fill:'none',stroke:state.series.color,'stroke-width':state.series.dash?'2.2':'2.7','stroke-dasharray':state.series.dash?'7 4':'none'}));state.accepted.filter(a=>a.age>=domain.x0&&a.age<=domain.x1).forEach(a=>{const offset=state.series.order*7;f.svg.appendChild(svgNode('line',{x1:f.x(a.age),y1:f.m.t+f.ph+2+offset,x2:f.x(a.age),y2:f.m.t+f.ph+7+offset,stroke:state.series.color,'stroke-width':'.8',opacity:'.6'}));});state.peaks.forEach((p,i)=>{const active=selectedState?.state===state&&selectedState.index===i,c=svgNode('circle',{cx:f.x(p.age),cy:f.y(p.d),r:active?4.5:3,fill:state.series.color,stroke:active?'#fff':'none','stroke-width':'1.5'});c.appendChild(svgNode('title',{},`${state.series.label}: пик ${p.age.toFixed(0)} млн лет; границы ${p.lo.toFixed(0)}–${p.hi.toFixed(0)}`));f.svg.appendChild(c);});});f.svg.appendChild(svgNode('text',{x:f.m.l+f.pw/2,y:f.height-8,'text-anchor':'middle',fill:'#465149','font-size':'10'},'Возраст, млн лет'));f.svg.addEventListener('click',event=>{const r=f.svg.getBoundingClientRect(),px=(event.clientX-r.left)/r.width*f.width,py=(event.clientY-r.top)/r.height*f.height,candidates=states.flatMap(state=>state.peaks.map((p,index)=>({state,index,distance:Math.hypot(f.x(p.age)-px,f.y(p.d)-py)})));if(px<f.m.l||px>f.m.l+f.pw||!candidates.length)return;const hit=candidates.reduce((a,b)=>a.distance<b.distance?a:b);geoSelections.set(dataset.id,{seriesId:hit.state.series.id,age:hit.state.peaks[hit.index].age});renderGeochronology();});}
function geoPeakButtons(container,dataset,states,selectedState,multiple){container.innerHTML=states.flatMap(state=>state.peaks.map((p,i)=>{const active=selectedState?.state===state&&selectedState.index===i;return `<button class="geo-peak ${active?'active':''}" data-series="${esc(state.series.id)}" data-peak="${i}" aria-pressed="${active}">${multiple?esc(state.series.label)+' · ':''}${p.age.toFixed(0)}</button>`;})).join('')||'Локальные пики не выделены';container.querySelectorAll('button').forEach(button=>button.onclick=()=>{const state=states.find(s=>s.series.id===button.dataset.series),i=Number(button.dataset.peak);geoSelections.set(dataset.id,{seriesId:state.series.id,age:state.peaks[i].age});renderGeochronology();});}
function renderGeoDataset(dataset,card){const h=Number(document.getElementById('geoBandwidth').value),domain=geoDomain(),allSeries=geoSeries(dataset),visibleSeries=geoVisibleSeries(dataset),states=geoStates(visibleSeries,domain,h),previous=geoSelections.get(dataset.id),selectedState=geoSelected(states,previous),chart=card.querySelector('.geo-chart'),legend=card.querySelector('.geo-series-legend'),chips=card.querySelector('.geo-peaks-full'),stats=card.querySelector('.geo-stats'),multiple=allSeries.length>1;legend.innerHTML=multiple?allSeries.map(s=>{const checked=geoVisibility.get(dataset.id).has(s.id);return `<label><input type="checkbox" data-series="${esc(s.id)}" ${checked?'checked':''}><i style="background:${s.color}"></i>${esc(s.label)} · ${s.analyses.filter(a=>a.accepted).length}/${s.analyses.length}</label>`;}).join(''):allSeries.map(s=>`<span><i style="background:${s.color}"></i>${esc(s.label)} · ${s.analyses.filter(a=>a.accepted).length}/${s.analyses.length}</span>`).join('');legend.querySelectorAll('input').forEach(input=>input.onchange=()=>{const visible=geoVisibility.get(dataset.id);if(input.checked)visible.add(input.dataset.series);else if(visible.size>1)visible.delete(input.dataset.series);else input.checked=true;if(!visible.has(geoSelections.get(dataset.id)?.seriesId))geoSelections.delete(dataset.id);renderGeochronology();});geoChart(dataset,chart,states,selectedState,domain,h,285);geoPeakButtons(chips,dataset,states,selectedState,multiple);const youngButton=card.querySelector('.geo-young-open');if(youngButton)youngButton.onclick=openYoungAnalysis;if(!selectedState){stats.textContent='Выберите пик на графике или кнопку под ним, чтобы рассчитать параметры анализов между соседними минимумами KDE.';return;}const state=selectedState.state,p=state.peaks[selectedState.index],subset=state.accepted.filter(a=>a.age>=p.lo&&a.age<=p.hi),wm=geoWeighted(subset),interval=`${p.lo.toFixed(0)}–${p.hi.toFixed(0)} млн лет`,prefix=`${state.series.label} · `;if(!wm){stats.textContent=`${prefix}пик ${p.age.toFixed(0)} млн лет · интервал ${interval} · анализов нет`;return;}if(wm.n===1){stats.textContent=`${prefix}пик ${p.age.toFixed(0)} млн лет · интервал ${interval} · n = 1 · ${wm.mean.toFixed(1)} ± ${wm.u2.toFixed(1)} млн лет (2σ). Для MSWD нужно ≥ 2 анализа.`;return;}stats.innerHTML=`<b>${esc(prefix)}пик ${p.age.toFixed(0)} млн лет</b> · интервал ${interval} · n = ${wm.n}<br>Взвешенное среднее: <b>${wm.mean.toFixed(1)} ± ${wm.u2.toFixed(1)} млн лет</b> (2σ внутр.)${wm.mswd>1?`; расширенное 2σ = ${wm.expanded.toFixed(1)}`:''}; MSWD = ${wm.mswd.toFixed(2)}; p(χ²) = ${wm.p<.001?'&lt;0,001':wm.p.toFixed(3).replace('.',',')}`;}
function renderGeochronology(){const grid=document.getElementById('geoGrid'),h=Number(document.getElementById('geoBandwidth').value),datasets=GEOCHRONOLOGY.datasets.filter(d=>d.kde_enabled!==false);document.getElementById('geoBandwidthLabel').textContent=`${h} млн лет`;grid.innerHTML=datasets.map(d=>{const series=geoSeries(d),counts=series.map(s=>`${esc(s.label)}: ${s.analyses.filter(a=>a.accepted).length}/${s.analyses.length}`).join('<br>'),sources=[...new Set(series.map(s=>s.source))];return `<article class="geo-card ${d.id===geoFocus?'focused':''}" data-dataset="${esc(d.id)}"><div class="geo-card-head"><div><strong>${esc(d.label)}</strong><span>${esc(d.subtitle)}</span><br><span><b>Объект:</b> ${esc(d.object)}</span></div><span>${counts}</span></div><div class="geo-series-legend"></div><div class="geo-chart-label">KDE · весь спектр <span>h = ${h} млн лет</span></div><div class="geo-chart"></div><div class="geo-peaks geo-peaks-full"></div>${d.id==='GUBINSKAYA-COMPARISON'?'<button class="geo-young-open">Анализ молодых пиков 700–900 млн лет</button>':''}<div class="geo-stats"></div>${d.published_result?`<div class="geo-source"><b>Опубликовано:</b> ${esc(d.published_result)}</div>`:''}${d.interpretation?`<div class="geo-source"><b>Проверка:</b> ${esc(d.interpretation)}</div>`:''}<div class="geo-source">Источник: ${sources.map(esc).join('; ')}${d.doi?` · <a href="${esc(d.doi)}" target="_blank" rel="noopener">DOI</a>`:''}${d.comparison?.doi?` · <a href="${esc(d.comparison.doi)}" target="_blank" rel="noopener">DOI Кочнева</a>`:''}${d.figure?` · <a href="${esc(d.figure)}" target="_blank">U–Pb диаграммы</a>`:''}${d.audit?` · <a href="${esc(d.audit)}" download>Аудит CSV</a>`:''}${d.report?` · <a href="${esc(d.report)}" target="_blank">Вывод</a>`:''}</div></article>`;}).join('');datasets.forEach(d=>{const card=grid.querySelector(`[data-dataset="${CSS.escape(d.id)}"]`);if(card)renderGeoDataset(d,card);});if(geoFocus){requestAnimationFrame(()=>grid.querySelector(`[data-dataset="${CSS.escape(geoFocus)}"]`)?.scrollIntoView({block:'center'}));}}
function openGeochronology(focusId=null){geoFocus=focusId;const modal=document.getElementById('geoModal');modal.classList.add('visible');modal.setAttribute('aria-hidden','false');requestAnimationFrame(renderGeochronology);}
function closeGeochronology(){const modal=document.getElementById('geoModal');modal.classList.remove('visible');modal.setAttribute('aria-hidden','true');geoFocus=null;}
let youngSeriesId='KOCHNEV-41812';
function renderYoungAnalysis(){const dataset=GEOCHRONOLOGY.datasets.find(d=>d.id==='GUBINSKAYA-COMPARISON');if(!dataset)return;const seriesList=geoSeries(dataset),series=seriesList.find(s=>s.id===youngSeriesId)||seriesList[0],accepted=series.analyses.filter(a=>a.accepted&&Number.isFinite(a.age)),ages=accepted.map(a=>a.age),rangeAges=ages.filter(a=>a>=700&&a<=900),h=15,x0=680,x1=920,steps=1200,xs=Array.from({length:steps+1},(_,i)=>x0+(x1-x0)*i/steps),dens=geoKdeValues(ages,xs,h),localPeaks=geoFeatures(xs,dens),regions=[{lo:700,hi:750,color:'rgba(145,91,45,.20)'},{lo:750,hi:850,color:'rgba(77,111,145,.22)'},{lo:850,hi:900,color:'rgba(73,121,80,.22)'}];regions.forEach((r,ri)=>{const candidates=localPeaks.filter(p=>p.age>=r.lo&&p.age<=r.hi),indices=xs.map((x,i)=>({x,i})).filter(q=>q.x>=r.lo&&q.x<=r.hi),fallback=indices.reduce((a,b)=>dens[a.i]>dens[b.i]?a:b),best=candidates.length?candidates.reduce((a,b)=>a.d>b.d?a:b):{age:fallback.x,d:dens[fallback.i]},members=rangeAges.filter(a=>a>=r.lo&&(ri===regions.length-1?a<=r.hi:a<r.hi));r.peak=best.age;r.d=best.d;r.n=members.length;r.pct=rangeAges.length?100*members.length/rangeAges.length:0;});const modal=document.getElementById('youngModal'),style=getComputedStyle(modal),bg=style.getPropertyValue('--ym-panel').trim(),textColor=style.getPropertyValue('--ym-text').trim(),muted=style.getPropertyValue('--ym-muted').trim(),grid=style.getPropertyValue('--ym-grid').trim(),container=document.getElementById('youngChart'),width=Math.max(760,container.clientWidth||1100),height=430,m={l:70,r:28,t:48,b:66},pw=width-m.l-m.r,ph=height-m.t-m.b,maxD=Math.max(...dens)*1.13,x=v=>m.l+(v-x0)/(x1-x0)*pw,y=v=>m.t+ph-v/maxD*ph,svg=svgNode('svg',{viewBox:`0 0 ${width} ${height}`,role:'img','aria-label':`${series.label}: KDE молодых возрастов 700–900 млн лет, h 15 млн лет`});container.innerHTML='';container.appendChild(svg);svg.appendChild(svgNode('rect',{x:0,y:0,width,height,fill:bg}));regions.forEach(r=>{svg.appendChild(svgNode('rect',{x:x(r.lo),y:m.t,width:x(r.hi)-x(r.lo),height:ph,fill:r.color}));svg.appendChild(svgNode('line',{x1:x(r.lo),y1:m.t,x2:x(r.lo),y2:m.t+ph,stroke:grid,'stroke-dasharray':'5 5'}));svg.appendChild(svgNode('text',{x:(x(r.lo)+x(r.hi))/2,y:m.t+24,'text-anchor':'middle',fill:textColor,'font-size':'15'},r.n?`≈${r.peak.toFixed(0)} млн лет`:'нет зерен'));svg.appendChild(svgNode('text',{x:(x(r.lo)+x(r.hi))/2,y:m.t+45,'text-anchor':'middle',fill:muted,'font-size':'12'},`n = ${r.n}; ${r.pct.toFixed(1)}%`));});svg.appendChild(svgNode('line',{x1:x(900),y1:m.t,x2:x(900),y2:m.t+ph,stroke:grid,'stroke-dasharray':'5 5'}));[700,750,800,850,900].forEach(t=>{svg.appendChild(svgNode('line',{x1:x(t),y1:m.t,x2:x(t),y2:m.t+ph,stroke:grid,'stroke-width':'1'}));svg.appendChild(svgNode('text',{x:x(t),y:m.t+ph+38,'text-anchor':'middle',fill:muted,'font-size':'12'},String(t)));});let path='';xs.forEach((v,i)=>{path+=(i?'L':'M')+x(v).toFixed(2)+','+y(dens[i]).toFixed(2);});svg.appendChild(svgNode('path',{d:path,fill:'none',stroke:'#76b9ee','stroke-width':'3'}));rangeAges.forEach(age=>svg.appendChild(svgNode('line',{x1:x(age),y1:m.t+ph+4,x2:x(age),y2:m.t+ph+18,stroke:muted,'stroke-width':'1'})));svg.appendChild(svgNode('line',{x1:m.l,y1:m.t+ph,x2:m.l+pw,y2:m.t+ph,stroke:grid}));svg.appendChild(svgNode('line',{x1:m.l,y1:m.t,x2:m.l,y2:m.t+ph,stroke:grid}));svg.appendChild(svgNode('text',{x:m.l+pw/2,y:height-10,'text-anchor':'middle',fill:muted,'font-size':'12'},'Best age, млн лет'));document.getElementById('youngMeta').innerHTML=`<span>${esc(series.label)}</span><span>700–900 млн лет: <b>${rangeAges.length}</b> из ${accepted.length} принятых анализов</span><span>KDE bandwidth = ${h} млн лет</span><span>доли внутри диапазона 700–900 млн лет</span>`;const buttons=document.getElementById('youngSeriesButtons');buttons.innerHTML=seriesList.map(s=>`<button data-young-series="${esc(s.id)}" class="${s.id===series.id?'active':''}">${esc(s.label)}</button>`).join('');buttons.querySelectorAll('button').forEach(button=>button.onclick=()=>{youngSeriesId=button.dataset.youngSeries;renderYoungAnalysis();});}
function openYoungAnalysis(){const modal=document.getElementById('youngModal');modal.classList.add('visible');modal.setAttribute('aria-hidden','false');youngSeriesId='KOCHNEV-41812';requestAnimationFrame(renderYoungAnalysis);}
function closeYoungAnalysis(){const modal=document.getElementById('youngModal');modal.classList.remove('visible');modal.setAttribute('aria-hidden','true');}
function project(lat,lon,z){const scale=256*2**z;const x=(lon+180)/360*scale;const sin=Math.sin(lat*Math.PI/180);const y=(.5-Math.log((1+sin)/(1-sin))/(4*Math.PI))*scale;return{x,y};}
function unproject(x,y,z){const scale=256*2**z;const lon=x/scale*360-180;const n=Math.PI-2*Math.PI*y/scale;const lat=180/Math.PI*Math.atan(Math.sinh(n));return{lat,lon};}
function screenPoint(lat,lon){const c=project(center.lat,center.lon,zoom),p=project(lat,lon,zoom);return{x:canvas.width/2+(p.x-c.x)*dpr,y:canvas.height/2+(p.y-c.y)*dpr};}
function getTile(z,x,y){const max=2**z,xw=((x%max)+max)%max;if(y<0||y>=max)return null;const key=`${z}/${xw}/${y}`;let e=tileCache.get(key);if(!e){const img=new Image();img.crossOrigin='anonymous';e={img,loaded:false,failed:false};img.onload=()=>{e.loaded=true;draw();};img.onerror=()=>{e.failed=true;draw();};img.src=tileTemplate.replace('{z}',z).replace('{x}',xw).replace('{y}',y);tileCache.set(key,e);}return e;}
function drawTiles(){const tileZoom=Math.max(0,Math.min(19,Math.floor(zoom))),factor=2**(zoom-tileZoom),c=project(center.lat,center.lon,tileZoom),viewW=canvas.width/(dpr*factor),viewH=canvas.height/(dpr*factor),left=c.x-viewW/2,top=c.y-viewH/2,x0=Math.floor(left/256),y0=Math.floor(top/256),x1=Math.ceil((left+viewW)/256),y1=Math.ceil((top+viewH)/256),tileSize=256*dpr*factor;ctx.fillStyle='#1c2828';ctx.fillRect(0,0,canvas.width,canvas.height);for(let x=x0;x<=x1;x++)for(let y=y0;y<=y1;y++){const e=getTile(tileZoom,x,y),sx=(x*256-left)*dpr*factor,sy=(y*256-top)*dpr*factor;if(e?.loaded)ctx.drawImage(e.img,sx,sy,tileSize+.5,tileSize+.5);else{ctx.fillStyle='#25302e';ctx.fillRect(sx,sy,tileSize+1,tileSize+1);}}}
const rasterTileCache=new Map();
let rasterFrame=0,rasterEnabled=!document.documentElement.classList.contains('mapshot'),rasterTransparency=35;
try{const saved=JSON.parse(localStorage.getItem('igarka-map-overlay')||'null');if(saved&&typeof saved.enabled==='boolean'&&typeof saved.transparency==='number'&&Number.isFinite(saved.transparency)){rasterEnabled=saved.enabled&&!document.documentElement.classList.contains('mapshot');rasterTransparency=Math.max(0,Math.min(100,saved.transparency));}}catch(e){}
function requestRasterDraw(){if(!rasterFrame)rasterFrame=requestAnimationFrame(()=>{rasterFrame=0;draw();});}
function rasterTile(level,x,y){const key=`${level}/${x}/${y}`;let entry=rasterTileCache.get(key);if(entry){rasterTileCache.delete(key);rasterTileCache.set(key,entry);return entry;}const img=new Image();entry={img,loaded:false,failed:false};img.onload=()=>{entry.loaded=true;requestRasterDraw();};img.onerror=()=>{entry.failed=true;document.getElementById('mapOverlayStatus').textContent='Часть карты недоступна: проверьте файлы слоя.';requestRasterDraw();};rasterTileCache.set(key,entry);img.src=RASTER_OVERLAY.template.replace('{level}',level).replace('{x}',x).replace('{y}',y)+'?v='+RASTER_OVERLAY.source_stamp.mtime_ns;while(rasterTileCache.size>64)rasterTileCache.delete(rasterTileCache.keys().next().value);return entry;}
function drawRasterOverlay(){if(!RASTER_OVERLAY||!rasterEnabled||rasterTransparency>=100)return;const b=RASTER_OVERLAY.bounds,nw=screenPoint(b[3],b[0]),se=screenPoint(b[1],b[2]),scaleX=(se.x-nw.x)/RASTER_OVERLAY.width,scaleY=(se.y-nw.y)/RASTER_OVERLAY.height;if(se.x<0||se.y<0||nw.x>canvas.width||nw.y>canvas.height)return;const level=Math.max(0,Math.min(RASTER_OVERLAY.max_level,Math.floor(Math.log2(1/Math.max(scaleX,scaleY))))),factor=2**level,info=RASTER_OVERLAY.levels[level],size=RASTER_OVERLAY.tile_size,tileW=size*factor*scaleX,tileH=size*factor*scaleY,x0=Math.max(0,Math.floor(-nw.x/tileW)),y0=Math.max(0,Math.floor(-nw.y/tileH)),x1=Math.min(Math.ceil(info.width/size)-1,Math.floor((canvas.width-nw.x)/tileW)),y1=Math.min(Math.ceil(info.height/size)-1,Math.floor((canvas.height-nw.y)/tileH));ctx.save();ctx.globalAlpha=1-rasterTransparency/100;ctx.imageSmoothingEnabled=true;ctx.beginPath();ctx.rect(nw.x,nw.y,se.x-nw.x,se.y-nw.y);ctx.clip();for(let y=y0;y<=y1;y++)for(let x=x0;x<=x1;x++){const entry=rasterTile(level,x,y);if(entry.loaded)ctx.drawImage(entry.img,nw.x+x*tileW,nw.y+y*tileH,entry.img.width*factor*scaleX,entry.img.height*factor*scaleY);}ctx.restore();}
function syncRasterControls(){document.getElementById('mapOverlayToggle').checked=rasterEnabled&&!!RASTER_OVERLAY;document.getElementById('mapOverlayToggle').disabled=!RASTER_OVERLAY;document.getElementById('mapOverlayOpacity').value=rasterTransparency;document.getElementById('mapOverlayOpacity').disabled=!rasterEnabled||!RASTER_OVERLAY;document.getElementById('mapOverlaySlider').classList.toggle('inactive',!rasterEnabled||!RASTER_OVERLAY);document.getElementById('mapOverlayValue').textContent=rasterTransparency+'%';if(!RASTER_OVERLAY)document.getElementById('mapOverlayStatus').textContent='Карта не подготовлена';}
function saveRasterSettings(){try{localStorage.setItem('igarka-map-overlay',JSON.stringify({enabled:rasterEnabled,transparency:rasterTransparency}));}catch(e){}}
document.getElementById('mapOverlayToggle').onchange=e=>{rasterEnabled=e.target.checked;syncRasterControls();saveRasterSettings();draw();};
document.getElementById('mapOverlayOpacity').oninput=e=>{rasterTransparency=Number(e.target.value);syncRasterControls();saveRasterSettings();draw();};
function positionRasterControls(){const toolbar=document.querySelector('.map-toolbar');document.querySelector('.map-layer-control').style.top=(toolbar.offsetTop+toolbar.offsetHeight+8)+'px';}
if(typeof ResizeObserver!=='undefined')new ResizeObserver(positionRasterControls).observe(document.querySelector('.map-toolbar'));else window.addEventListener('resize',positionRasterControls);
syncRasterControls();positionRasterControls();
function pointLabel(pointId){return String(pointId||'');}
function drawMapPin(p,color,isActive,size=8.5){const s=size*dpr;ctx.save();ctx.translate(p.x,p.y);ctx.shadowColor='rgba(0,0,0,.48)';ctx.shadowBlur=4*dpr;ctx.fillStyle=color;ctx.strokeStyle=isActive?'#ffd45f':'rgba(255,255,255,.96)';ctx.lineWidth=(isActive?2.5:1.3)*dpr;ctx.beginPath();ctx.moveTo(-.62*s,-.75*s);ctx.lineTo(0,0);ctx.lineTo(.62*s,-.75*s);ctx.closePath();ctx.fill();ctx.stroke();ctx.beginPath();ctx.arc(0,-1.42*s,s,0,Math.PI*2);ctx.fill();ctx.stroke();ctx.shadowBlur=0;ctx.beginPath();ctx.arc(0,-1.42*s,.30*s,0,Math.PI*2);ctx.fillStyle='white';ctx.fill();ctx.restore();}
function drawMapLabel(p,label,side='right',offsetY=0){if(!showMapLabels)return;ctx.save();const labelSize=document.documentElement.classList.contains('mapshot')?22:10;ctx.font=`700 ${labelSize*dpr}px -apple-system,BlinkMacSystemFont,"Segoe UI",Arial`;const textW=ctx.measureText(label).width,x=side==='left'?p.x-12*dpr-textW:p.x+12*dpr,y=p.y-15*dpr+offsetY*dpr;ctx.textAlign='left';ctx.textBaseline='middle';ctx.lineJoin='round';ctx.lineWidth=(document.documentElement.classList.contains('mapshot')?5:3)*dpr;ctx.strokeStyle='rgba(18,25,21,.92)';ctx.strokeText(label,x,y);ctx.fillStyle='white';ctx.fillText(label,x,y);ctx.restore();}
function draw(){
  if(!canvas.width)return;
  canvas.dataset.zoom=zoom.toFixed(3);ctx.setTransform(1,0,0,1,0,0);ctx.clearRect(0,0,canvas.width,canvas.height);drawTiles();drawRasterOverlay();hitItems.length=0;
  const groups=new Map();
  for(const s of filtered())if(Number.isFinite(s.lat)&&Number.isFinite(s.lon)){const key=`${s.lat.toFixed(6)},${s.lon.toFixed(6)}`;if(!groups.has(key))groups.set(key,[]);groups.get(key).push(s);}
  for(const group of groups.values()){
    const first=group[0],p=screenPoint(first.lat,first.lon),label=pointLabel(first.point_id),isActive=group.some(s=>s.sample===activeSample);
    if(p.x<-55||p.y<-30||p.x>canvas.width+55||p.y>canvas.height+30)continue;
    drawMapPin(p,first.year===2026?'#c47725':'#19765f',isActive,8.5);drawMapLabel(p,label);
    hitItems.push({kind:'sample',x:p.x,y:p.y-11*dpr,r:13*dpr,group});
  }
  if(showLiteraturePoints)for(const point of literaturePoints){
    const p=screenPoint(point.lat,point.lon),isActive=point.id===activeLiterature;
    if(p.x<-70||p.y<-40||p.x>canvas.width+70||p.y>canvas.height+40)continue;
    drawMapPin(p,'#6b5a92',isActive,7.5);drawMapLabel(p,point.sample,point.sample==='1812-1'?'left':'right',point.sample==='1812-2'?8:0);
    hitItems.push({kind:'literature',x:p.x,y:p.y-10*dpr,r:12*dpr,point});
  }
}
function fitTo(arr){const pts=arr.filter(s=>Number.isFinite(s.lat)&&Number.isFinite(s.lon));if(!pts.length)return;if(zoomAnimation)cancelAnimationFrame(zoomAnimation);zoomAnimation=0;const minLat=Math.min(...pts.map(s=>s.lat)),maxLat=Math.max(...pts.map(s=>s.lat)),minLon=Math.min(...pts.map(s=>s.lon)),maxLon=Math.max(...pts.map(s=>s.lon));center={lat:(minLat+maxLat)/2,lon:(minLon+maxLon)/2};for(let z=17;z>=2;z--){const a=project(maxLat,minLon,z),b=project(minLat,maxLon,z);if(Math.abs(b.x-a.x)*dpr<canvas.width*.78&&Math.abs(b.y-a.y)*dpr<canvas.height*.78){zoom=z;break;}}targetZoom=zoom;draw();}
function visibleLiterature(){return showLiteraturePoints?literaturePoints:[];}
function fitVisible(arr){fitTo([...arr,...visibleLiterature()]);}
function syncMapToggles(){const literatureButton=document.getElementById('literatureToggle'),labelsButton=document.getElementById('labelsToggle'),legend=document.getElementById('kochnevLegend');literatureButton.setAttribute('aria-pressed',String(showLiteraturePoints));labelsButton.setAttribute('aria-pressed',String(showMapLabels));if(legend)legend.style.opacity=showLiteraturePoints?'1':'.42';}
function resize(){const r=canvas.getBoundingClientRect();canvas.width=Math.max(1,Math.round(r.width*dpr));canvas.height=Math.max(1,Math.round(r.height*dpr));draw();}
function showPopup(group,x,y){const point=group[0];popup.style.left=Math.min(Math.max(x/dpr+12,8),canvas.clientWidth-280)+'px';popup.style.top=Math.min(Math.max(y/dpr+12,8),canvas.clientHeight-190)+'px';popup.innerHTML=`<div class="popup-title">${esc(point.point_id)} · ${group.length} ${group.length===1?'образец':group.length<5?'образца':'образцов'}</div>${point.point_description?`<div class="popup-point">${esc(point.point_description)}</div>`:''}`+group.map(s=>`<button class="${s.sample===activeSample?'active':''}" data-sample="${esc(s.sample)}"><b>${esc(s.sample)}</b><br><span style="color:#68736c">${esc(s.location||'Без описания')}</span></button>`).join('');popup.classList.add('visible');popup.querySelectorAll('button').forEach(btn=>btn.onclick=()=>openSample(btn.dataset.sample,false));}
function showLiteraturePopup(points,x,y){popup.style.left=Math.min(Math.max(x/dpr+12,8),canvas.clientWidth-280)+'px';popup.style.top=Math.min(Math.max(y/dpr+12,8),canvas.clientHeight-190)+'px';popup.innerHTML=`<div class="popup-title">Точки Кочнева · ${points.length}</div>`+points.map(p=>`<button class="${p.id===activeLiterature?'active':''}" data-literature="${esc(p.id)}"><b>${esc(p.sample)}</b><br><span style="color:#68736c">${esc(p.formation)}</span></button>`).join('');popup.classList.add('visible');popup.querySelectorAll('button').forEach(btn=>btn.onclick=()=>{popup.classList.remove('visible');openLiterature(btn.dataset.literature,false);});}
function showMapChoice(hits,x,y){const entries=[],seen=new Set();hits.forEach(hit=>{if(hit.kind==='sample')hit.group.forEach(s=>{const key=`sample:${s.sample}`;if(!seen.has(key)){seen.add(key);entries.push({kind:'sample',id:s.sample,title:s.sample,subtitle:s.location||s.description||'Наш образец',source:'Малышев, Пасенко'});}});else{const p=hit.point,key=`literature:${p.id}`;if(!seen.has(key)){seen.add(key);entries.push({kind:'literature',id:p.id,title:p.sample,subtitle:p.formation||p.material||'Точка Кочнева',source:'Кочнев'});}}});popup.style.left=Math.min(Math.max(x/dpr+12,8),canvas.clientWidth-300)+'px';popup.style.top=Math.min(Math.max(y/dpr+12,8),canvas.clientHeight-210)+'px';popup.innerHTML=`<div class="popup-title">Выберите точку · ${entries.length}</div>`+entries.map(e=>`<button data-kind="${e.kind}" data-id="${esc(e.id)}"><b>${esc(e.title)}</b> <span style="color:#68736c">· ${esc(e.source)}</span><br><span style="color:#68736c">${esc(e.subtitle)}</span></button>`).join('');popup.classList.add('visible');popup.querySelectorAll('button').forEach(btn=>btn.onclick=()=>{popup.classList.remove('visible');if(btn.dataset.kind==='sample')openSample(btn.dataset.id,false);else openLiterature(btn.dataset.id,false);});}
function screenToLatLon(sx,sy){const c=project(center.lat,center.lon,zoom);return unproject(c.x+(sx-canvas.width/2)/dpr,c.y+(sy-canvas.height/2)/dpr,zoom);}
function positionAnchor(anchor,sx,sy,newZoom){zoom=newZoom;const p=project(anchor.lat,anchor.lon,zoom),c={x:p.x-(sx-canvas.width/2)/dpr,y:p.y-(sy-canvas.height/2)/dpr};center=unproject(c.x,c.y,zoom);}
function smoothZoom(delta,sx=canvas.width/2,sy=canvas.height/2){const anchor=screenToLatLon(sx,sy),startZoom=zoom,startTime=performance.now();targetZoom=Math.max(3,Math.min(19,(zoomAnimation?targetZoom:zoom)+delta));const finishZoom=targetZoom;if(zoomAnimation)cancelAnimationFrame(zoomAnimation);const frame=now=>{const t=Math.min(1,(now-startTime)/220),ease=1-(1-t)**3;positionAnchor(anchor,sx,sy,startZoom+(finishZoom-startZoom)*ease);draw();if(t<1)zoomAnimation=requestAnimationFrame(frame);else{zoomAnimation=0;targetZoom=zoom;}};zoomAnimation=requestAnimationFrame(frame);}
function gestureZoom(delta,sx,sy){if(zoomAnimation)cancelAnimationFrame(zoomAnimation);zoomAnimation=0;targetZoom=zoom;const anchor=screenToLatLon(sx,sy),next=Math.max(3,Math.min(19,zoom+delta));if(next===zoom)return;positionAnchor(anchor,sx,sy,next);targetZoom=zoom;draw();}
canvas.addEventListener('wheel',e=>{e.preventDefault();const rect=canvas.getBoundingClientRect(),sx=(e.clientX-rect.left)*dpr,sy=(e.clientY-rect.top)*dpr,units=e.deltaMode===1?e.deltaY*16:e.deltaMode===2?e.deltaY*120:e.deltaY,sensitivity=e.ctrlKey ? .045 : .01,delta=Math.max(-.45,Math.min(.45,-units*sensitivity));if(Math.abs(delta)<.001)return;popup.classList.remove('visible');gestureZoom(delta,sx,sy);},{passive:false});
canvas.addEventListener('pointerdown',e=>{if(zoomAnimation)cancelAnimationFrame(zoomAnimation);zoomAnimation=0;targetZoom=zoom;dragging=true;moved=false;const c=project(center.lat,center.lon,zoom);dragStart={x:e.clientX,y:e.clientY,cx:c.x,cy:c.y};canvas.classList.add('dragging');canvas.setPointerCapture(e.pointerId);});
canvas.addEventListener('pointermove',e=>{if(!dragging)return;const dx=e.clientX-dragStart.x,dy=e.clientY-dragStart.y;if(Math.hypot(dx,dy)>3)moved=true;center=unproject(dragStart.cx-dx,dragStart.cy-dy,zoom);popup.classList.remove('visible');draw();});
canvas.addEventListener('pointerup',e=>{dragging=false;canvas.classList.remove('dragging');if(moved)return;const rect=canvas.getBoundingClientRect(),x=(e.clientX-rect.left)*dpr,y=(e.clientY-rect.top)*dpr,hit=[...hitItems].reverse().find(h=>Math.hypot(h.x-x,h.y-y)<=h.r);if(!hit){popup.classList.remove('visible');return;}const nearby=hitItems.filter(h=>Math.hypot(h.x-hit.x,h.y-hit.y)<=36*dpr),choiceCount=nearby.reduce((n,h)=>n+(h.kind==='sample'?h.group.length:1),0);if(choiceCount>1){showMapChoice(nearby,x,y);return;}if(hit.kind==='literature'){popup.classList.remove('visible');openLiterature(hit.point.id,false);}else{showPopup(hit.group,x,y);openSample(hit.group[0].sample,false);}});
document.getElementById('zoomIn').onclick=()=>smoothZoom(1);document.getElementById('zoomOut').onclick=()=>smoothZoom(-1);document.getElementById('fit').onclick=()=>fitVisible(filtered());
document.getElementById('literatureToggle').onclick=()=>{showLiteraturePoints=!showLiteraturePoints;popup.classList.remove('visible');if(!showLiteraturePoints)activeLiterature=null;syncMapToggles();draw();};
document.getElementById('labelsToggle').onclick=()=>{showMapLabels=!showMapLabels;syncMapToggles();draw();};
document.getElementById('workflowOpen').onclick=openWorkflow;document.getElementById('workflowClose').onclick=closeWorkflow;['workflowSearch','workflowYear','workflowFilter'].forEach(id=>document.getElementById(id).addEventListener('input',renderWorkflowTable));
document.getElementById('chemOpen').onclick=()=>openChemistry(activeSample);document.getElementById('chemClose').onclick=closeChemistry;document.getElementById('chemCurrent').onclick=()=>{if(chemActive){chemSelected=new Set([chemActive]);renderChemSidebar();renderChemCharts();}};document.getElementById('chemAll').onclick=()=>{chemSelected=new Set(samples.filter(s=>s.geochem).map(s=>s.sample));if(!chemActive)chemActive=[...chemSelected][0];renderChemSidebar();renderChemCharts();};
document.getElementById('geoOpen').onclick=()=>openGeochronology();document.getElementById('geoClose').onclick=closeGeochronology;document.getElementById('geoBandwidth').oninput=renderGeochronology;
document.getElementById('youngClose').onclick=closeYoungAnalysis;document.getElementById('youngTheme').onclick=()=>{const modal=document.getElementById('youngModal'),light=modal.classList.toggle('light');document.getElementById('youngTheme').textContent=light?'Тёмная тема':'Светлая тема';renderYoungAnalysis();};
document.getElementById('upbClose').onclick=closeUpb;
document.getElementById('upbLimit').oninput=renderUpb;
document.getElementById('upbMethod').onchange=()=>{upbView=upbFitView(upbActiveRows());renderUpb();};
document.getElementById('upbContext').onchange=renderUpb;
document.getElementById('upbCompare').onclick=openUpbComparison;
document.querySelectorAll('.upb-preset').forEach(b=>b.onclick=()=>{document.getElementById('upbLimit').value=b.dataset.limit;renderUpb();});
document.getElementById('upbZoomIn').onclick=()=>upbScaleView(.67);
document.getElementById('upbZoomOut').onclick=()=>upbScaleView(1.5);
document.getElementById('upbZoomFit').onclick=()=>{upbView=upbFitView(upbActiveRows());renderUpbConcordia();};
document.getElementById('upbZoomReset').onclick=()=>{upbView=upbWholeView();renderUpbConcordia();};
const upbConcordiaBox=document.getElementById('upbConcordia');
upbConcordiaBox.addEventListener('wheel',e=>{e.preventDefault();const r=upbConcordiaBox.getBoundingClientRect(),cx=Math.max(0,Math.min(1,(e.clientX-r.left)/r.width)),cy=Math.max(0,Math.min(1,(e.clientY-r.top)/r.height));upbScaleView(Math.exp(e.deltaY*.0015),cx,cy);},{passive:false});
upbConcordiaBox.addEventListener('dblclick',e=>{const r=upbConcordiaBox.getBoundingClientRect();upbScaleView(.55,(e.clientX-r.left)/r.width,(e.clientY-r.top)/r.height);});
upbConcordiaBox.addEventListener('pointerdown',e=>{upbPointers.set(e.pointerId,{x:e.clientX,y:e.clientY});upbConcordiaBox.setPointerCapture(e.pointerId);const r=upbConcordiaBox.getBoundingClientRect();if(upbPointers.size===1)upbDrag={x:e.clientX,y:e.clientY,view:{...(upbView||upbWholeView())},w:r.width,h:r.height};else if(upbPointers.size===2){const p=[...upbPointers.values()],dx=p[0].x-p[1].x,dy=p[0].y-p[1].y;upbPinch={distance:Math.max(1,Math.hypot(dx,dy)),view:{...(upbView||upbWholeView())},rect:r};upbDrag=null;}});
upbConcordiaBox.addEventListener('pointermove',e=>{if(!upbPointers.has(e.pointerId))return;upbPointers.set(e.pointerId,{x:e.clientX,y:e.clientY});if(upbPointers.size===2&&upbPinch){const p=[...upbPointers.values()],dx=p[0].x-p[1].x,dy=p[0].y-p[1].y,dist=Math.max(1,Math.hypot(dx,dy)),mx=(p[0].x+p[1].x)/2,my=(p[0].y+p[1].y)/2,cx=(mx-upbPinch.rect.left)/upbPinch.rect.width,cy=(my-upbPinch.rect.top)/upbPinch.rect.height;upbView=upbZoomedView(upbPinch.view,upbPinch.distance/dist,cx,cy);renderUpbConcordia();}else if(upbPointers.size===1&&upbDrag){const v=upbDrag.view,dx=(e.clientX-upbDrag.x)/upbDrag.w*(v.x1-v.x0),dy=(e.clientY-upbDrag.y)/upbDrag.h*(v.y1-v.y0);upbView={x0:v.x0-dx,x1:v.x1-dx,y0:v.y0+dy,y1:v.y1+dy};renderUpbConcordia();}});
function endUpbPointer(e){upbPointers.delete(e.pointerId);upbPinch=null;if(upbPointers.size===1){const p=[...upbPointers.values()][0],r=upbConcordiaBox.getBoundingClientRect();upbDrag={x:p.x,y:p.y,view:{...(upbView||upbWholeView())},w:r.width,h:r.height};}else upbDrag=null;}
upbConcordiaBox.addEventListener('pointerup',endUpbPointer);upbConcordiaBox.addEventListener('pointercancel',endUpbPointer);
document.getElementById('lightClose').onclick=closeLightbox;document.getElementById('lightPrev').onclick=()=>stepLightbox(-1);document.getElementById('lightNext').onclick=()=>stepLightbox(1);document.getElementById('lightbox').addEventListener('click',e=>{if(e.target.id==='lightbox'||e.target.classList.contains('lightstage'))closeLightbox();});
['search','year','status'].forEach(id=>document.getElementById(id).addEventListener('input',()=>{popup.classList.remove('visible');renderList();fitVisible(filtered());}));
let chemResizeTimer=0,geoResizeTimer=0,youngResizeTimer=0,upbResizeTimer=0;window.addEventListener('resize',()=>{resize();if(document.getElementById('chemModal').classList.contains('visible')){clearTimeout(chemResizeTimer);chemResizeTimer=setTimeout(renderChemCharts,120);}if(document.getElementById('geoModal').classList.contains('visible')){clearTimeout(geoResizeTimer);geoResizeTimer=setTimeout(renderGeochronology,120);}if(document.getElementById('youngModal').classList.contains('visible')){clearTimeout(youngResizeTimer);youngResizeTimer=setTimeout(renderYoungAnalysis,120);}if(document.getElementById('upbModal').classList.contains('visible')){clearTimeout(upbResizeTimer);upbResizeTimer=setTimeout(renderUpb,120);}});document.addEventListener('keydown',e=>{if(e.key==='Escape'){popup.classList.remove('visible');closeWorkflow();closeLightbox();closeChemistry();closeYoungAnalysis();closeGeochronology();closeUpb();}if(document.getElementById('lightbox').classList.contains('visible')&&e.key==='ArrowLeft')stepLightbox(-1);if(document.getElementById('lightbox').classList.contains('visible')&&e.key==='ArrowRight')stepLightbox(1);});
function guideSectionForSample(id){if(['SM25-01','SM25-02','2007-15'].includes(id))return 'ludovskaya';if(['SM25-05','SM26-07','1812-2'].includes(id))return 'gubinskaya';if(['SM25-03','SM25-07','SM26-06','1812-1'].includes(id))return 'intrusions';if(/^SM25-(08|09|10|11|12|13|14|15|16)/.test(id)||['SM26-03','SM26-04'].includes(id))return 'igarskaya';return 'overview';}
function addGuideCardLink(){const detail=document.getElementById('detail');if(!detail||detail.querySelector('.guide-card-link'))return;const id=activeLiterature?literaturePoints.find(p=>p.id===activeLiterature)?.sample:activeSample;if(!id)return;const section=document.createElement('div');section.className='section guide-card-link';const link=document.createElement('a');link.href='guide.html#'+guideSectionForSample(id);link.textContent='Об объекте в путеводителе →';section.appendChild(link);detail.prepend(section);}
const guideButton=document.createElement('button');guideButton.id='guideOpen';guideButton.textContent='Путеводитель';guideButton.title='Геология Игарки: разрезы, образцы и источники';guideButton.onclick=()=>{location.href='guide.html';};document.querySelector('.map-toolbar').appendChild(guideButton);
new MutationObserver(addGuideCardLink).observe(document.getElementById('detail'),{childList:true});
function restorePortalLink(){const params=new URLSearchParams(location.search),sampleId=params.get('sample'),literatureId=params.get('literature'),datasetId=params.get('dataset'),view=params.get('view');if(sampleId&&samples.some(s=>s.sample===sampleId))openSample(sampleId,true);else if(literatureId&&showLiteraturePoints&&literaturePoints.some(p=>p.id===literatureId))openLiterature(literatureId,true);if(mapshotSamples.size&&!document.documentElement.classList.contains('mapshot')){showLiteraturePoints=false;syncMapToggles();fitVisible(filtered());}if(view==='upb'&&sampleId)openUpb(sampleId);else if(view==='upb-comparison')openUpbComparison();else if(datasetId&&GEOCHRONOLOGY.datasets.some(d=>d.id===datasetId))openGeochronology(datasetId);}
initControls();renderSummary();syncMapToggles();resize();renderList();fitVisible(filtered());if(samples.length)openSample(samples[0].sample,false);restorePortalLink();addGuideCardLink();
</script>
</body>
</html>'''
    return (
        template.replace("__PAYLOAD__", payload)
        .replace("__SPIDER_ELEMENTS__", spider_payload)
        .replace("__PRIMITIVE_MANTLE__", primitive_payload)
        .replace("__REFERENCES__", references_payload)
        .replace("__GEOCHRONOLOGY__", geochronology_payload)
        .replace("__LITERATURE_POINTS__", literature_points_payload)
        .replace("__RASTER_OVERLAY__", (ROOT / "map_overlay" / "metadata.json").read_text(encoding="utf-8") if (ROOT / "map_overlay" / "metadata.json").exists() else "null")
    )


def main() -> None:
    build_map_overlay()
    upb_environment = os.environ.copy()
    upb_environment.setdefault("MPLCONFIGDIR", "/private/tmp/igarka-mpl")
    subprocess.run([sys.executable, str(UPB_BUILDER)], check=True, env=upb_environment)
    records = load_records()
    photo_count, unmatched = attach_photos(records)
    thin_section_count, unmatched_thin_sections = attach_thin_sections(records)
    petrography_count, unmatched_petrography = attach_petrography(records)
    upb_count = attach_upb(records)
    geochemistry_count = attach_geochemistry(records)
    spider_elements, primitive_rows = read_geochemical_csv(PM_NORMALIZER)
    _, references = read_geochemical_csv(REFERENCE_CHEMISTRY)
    geochronology = json.loads(GEOCHRONOLOGY_DATA.read_text(encoding="utf-8"))
    primitive_mantle = primitive_rows[0]["values"]
    OUTPUT.write_text(
        build_html(records, spider_elements, primitive_mantle, references, geochronology), encoding="utf-8"
    )
    print(f"Created {build_guide()}")
    print(
        f"Created {OUTPUT} with {len(records)} records, {photo_count} outcrop photos "
        f"{thin_section_count} thin-section photos, {petrography_count} petrographic descriptions "
        f"{upb_count} U-Pb results and {geochemistry_count} geochemical datasets"
    )
    if unmatched:
        print("Unmatched outcrop photos: " + ", ".join(unmatched))
    if unmatched_thin_sections:
        print("Unmatched thin-section photos: " + ", ".join(unmatched_thin_sections))
    if unmatched_petrography:
        print("Unmatched petrographic descriptions: " + ", ".join(unmatched_petrography))


if __name__ == "__main__":
    main()
