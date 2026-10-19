#!/usr/bin/env python3
"""Horizontal TAS + PM-normalized figure for Igarka Formation tuffs."""

from __future__ import annotations

import csv
import math
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

from make_sm25_07_presentation_figure import TAS_FIELDS, TAS_LABELS


ROOT = Path(__file__).resolve().parent
DATA_ROOT = ROOT.parent
OUT_PNG = ROOT / "assets" / "Igarka_tuffs_TAS_multielement_horizontal_20cm.png"
OUT_PDF = ROOT / "assets" / "Igarka_tuffs_TAS_multielement_horizontal_20cm.pdf"

SAMPLES = ["SM25-08", "SM25-10*", "SM25-10/3", "SM25-11/2", "SM25-12/2", "SM25-12/3"]
TUFF_COLOR = "#126e64"

# 20 × 8.2 cm at 300 dpi.
WIDTH, HEIGHT = 2362, 969
FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def key(value: object) -> str:
    return "".join(str(value or "").upper().split())


def read_major() -> dict[str, dict[str, float]]:
    ws = load_workbook(DATA_ROOT / "Пасенко_Малышев_19_проб.xlsx", data_only=True, read_only=True)["Лист1"]
    headers = [str(cell.value or "").replace(" (%)", "").strip() for cell in ws[4]]
    result = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        sample = key(row[0])
        if sample in SAMPLES:
            result[sample] = {headers[i]: float(value) for i, value in enumerate(row) if i and value is not None}
    return result


def read_trace() -> dict[str, dict[str, float]]:
    ws = load_workbook(DATA_ROOT / "igarka_trace.xlsx", data_only=True, read_only=True)["Sheet1"]
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        sample = key(row[0])
        if sample in SAMPLES:
            result[sample] = {
                headers[i]: float(value)
                for i, value in enumerate(row)
                if i >= 2 and headers[i] and value is not None
            }
    return result


def read_csv(path: Path) -> tuple[list[str], dict[str, dict[str, float]]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    elements = [value.strip() for value in rows[0][1:]]
    datasets = {}
    for row in rows[1:]:
        if row and row[0].strip():
            datasets[row[0].strip()] = {
                element: float(value)
                for element, value in zip(elements, row[1:])
                if element and str(value).strip()
            }
    return elements, datasets


def marker(draw: ImageDraw.ImageDraw, x: float, y: float, shape: str, color: str, radius=10) -> None:
    r = radius
    if shape == "circle":
        draw.ellipse((x-r, y-r, x+r, y+r), fill=color, outline="white", width=3)
    elif shape == "square":
        draw.rectangle((x-r, y-r, x+r, y+r), fill=color, outline="white", width=3)
    elif shape == "diamond":
        draw.polygon([(x, y-r-2), (x+r+2, y), (x, y+r+2), (x-r-2, y)], fill=color, outline="white")
    elif shape == "triangle":
        draw.polygon([(x, y-r-3), (x+r+3, y+r), (x-r-3, y+r)], fill=color, outline="white")
    elif shape == "down":
        draw.polygon([(x-r-3, y-r), (x+r+3, y-r), (x, y+r+3)], fill=color, outline="white")
    else:
        points = [(x + r * math.cos(math.pi/3*i), y + r * math.sin(math.pi/3*i)) for i in range(6)]
        draw.polygon(points, fill=color, outline="white")


def dashed(draw, coords, fill, width=4, dash=18, gap=11):
    for (x1, y1), (x2, y2) in zip(coords, coords[1:]):
        dx, dy = x2 - x1, y2 - y1
        length = math.hypot(dx, dy)
        if length == 0:
            continue
        ux, uy = dx / length, dy / length
        position = 0.0
        while position < length:
            end = min(length, position + dash)
            draw.line((x1 + ux*position, y1 + uy*position, x1 + ux*end, y1 + uy*end), fill=fill, width=width)
            position += dash + gap


def main() -> None:
    majors = read_major()
    traces = read_trace()
    missing = [sample for sample in SAMPLES if sample not in majors or sample not in traces]
    if missing:
        raise RuntimeError("Missing geochemical data: " + ", ".join(missing))

    elements, pm_rows = read_csv(DATA_ROOT / "PM_1989.csv")
    _, refs = read_csv(DATA_ROOT / "morb.csv")
    pm = pm_rows["Primitive Mantle McDonough & Sun 1989"]

    image = Image.new("RGB", (WIDTH, HEIGHT), "white")
    draw = ImageDraw.Draw(image)
    sizes = (17, 18, 19, 20, 21, 22, 23, 24, 25, 27, 29, 31, 34, 36)
    regular = {size: ImageFont.truetype(FONT, size) for size in sizes}
    bold = {size: ImageFont.truetype(FONT_BOLD, size) for size in (25, 27, 29, 31, 34, 36)}

    def label(x, y, value, size=27, anchor="mm", color="#27312c", is_bold=False):
        draw.text((x, y), str(value), font=(bold if is_bold else regular)[size], fill=color, anchor=anchor)

    def multiline_label(x, y, lines, size, color="#3f4c46"):
        draw.multiline_text((x, y), "\n".join(lines), font=regular[size], fill=color, anchor="mm", align="center", spacing=1)

    def rotated_label(x, y, value, size=27, angle=90, color="#27312c"):
        font = regular[size]
        box = font.getbbox(str(value))
        tile = Image.new("RGBA", (box[2]-box[0]+12, box[3]-box[1]+12), (255,255,255,0))
        td = ImageDraw.Draw(tile)
        td.text((6-box[0], 6-box[1]), str(value), font=font, fill=color)
        tile = tile.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC)
        image.paste(tile, (round(x-tile.width/2), round(y-tile.height/2)), tile)

    # Left panel: TAS.
    left, right, top, bottom = 112, 1128, 88, 786
    xmap = lambda value: left + (value-38)/39*(right-left)
    ymap = lambda value: bottom - value/16*(bottom-top)
    label(left, 42, "а) TAS · туфы Игарской толщи", 36, "lm", is_bold=True)
    for tick in (0, 5, 10, 15):
        y = ymap(tick)
        draw.line((left, y, right, y), fill="#dfe5e2", width=2)
        label(left-16, y, tick, 24, "rm", "#4d5a54")
    for tick in (40, 50, 60, 70):
        x = xmap(tick)
        draw.line((x, top, x, bottom), fill="#edf0ee", width=2)
        label(x, bottom+29, tick, 24, "mm", "#4d5a54")
    for index, (_, points) in enumerate(TAS_FIELDS):
        coords = [(xmap(max(38,min(77,x))), ymap(max(0,min(16,y)))) for x,y in points]
        draw.polygon(coords, fill="#eef3f0" if index%2==0 else "#f8f9f8", outline="#87938d", width=2)
    for lx, ly, lines, size in TAS_LABELS:
        multiline_label(xmap(lx), ymap(ly), lines, size)
    draw.line((left,top,left,bottom), fill="#39443f", width=3)
    draw.line((left,bottom,right,bottom), fill="#39443f", width=3)
    label((left+right)/2, bottom+78, "SiO₂, мас. %", 29)
    rotated_label(37, (top+bottom)/2, "Na₂O + K₂O, мас. %", 29)

    oxides = ["Na2O","MgO","Al2O3","SiO2","P2O5","K2O","CaO","TiO2","MnO","Fe2O3"]
    for sample in SAMPLES:
        major = majors[sample]
        factor = 100/sum(major.get(name,0) for name in oxides)
        sio2 = major["SiO2"]*factor
        alkalis = (major["Na2O"]+major["K2O"])*factor
        marker(draw, xmap(sio2), ymap(alkalis), "circle", TUFF_COLOR, 10)

    # Right panel: primitive-mantle normalized patterns.
    mleft, mright, mtop, mbottom = 1270, 2325, 88, 786
    mx = lambda i: mleft + i/(len(elements)-1)*(mright-mleft)
    my = lambda value: mbottom - (math.log10(value)+1)/4*(mbottom-mtop)
    label(mleft, 42, "б) Мультиэлементная диаграмма", 36, "lm", is_bold=True)
    for value in (0.1,1,10,100,1000):
        y = my(value)
        draw.line((mleft,y,mright,y), fill="#dbe1de", width=2)
        label(mleft-16,y,f"{value:g}",24,"rm","#4d5a54")
    draw.line((mleft,mtop,mleft,mbottom), fill="#39443f", width=3)
    draw.line((mleft,mbottom,mright,mbottom), fill="#39443f", width=3)
    for i, element in enumerate(elements):
        rotated_label(mx(i), mbottom+45, element, 20, 90, "#39443f")
    label((mleft+mright)/2, bottom+78, "Элементы", 29)
    rotated_label(1194, (mtop+mbottom)/2, "Порода / примитивная мантия", 27)

    ref_series = [
        ("N-MORB", refs["N-MORB (Gale et al., 2013)"], "#536a93"),
        ("OIB", refs["OIB (Sun&McDoniugh, 1989)"], "#b06b2b"),
    ]
    for _, values, color in ref_series:
        coords = [(mx(i), my(values[element]/pm[element])) for i,element in enumerate(elements)]
        dashed(draw, coords, color, 4)

    for sample in SAMPLES:
        values = dict(traces[sample])
        major = majors[sample]
        values.update(K=major["K2O"]*8301.5, P=major["P2O5"]*4364.2, Ti=major["TiO2"]*5993.4)
        coords = [(mx(i), my(values[element]/pm[element])) for i,element in enumerate(elements)]
        draw.line(coords, fill=TUFF_COLOR, width=5, joint="curve")

    # Compact legend in the free upper-right part of the multielement panel.
    legend_x, legend_y = 1810, 126
    draw.line((legend_x, legend_y, legend_x+78, legend_y), fill=TUFF_COLOR, width=6)
    label(legend_x+94, legend_y, "туфы Игарской толщи", 23, "lm")
    dashed(draw, [(legend_x,legend_y+42),(legend_x+78,legend_y+42)], "#536a93", 4)
    label(legend_x+94, legend_y+42, "N-MORB", 23, "lm")
    dashed(draw, [(legend_x,legend_y+84),(legend_x+78,legend_y+84)], "#b06b2b", 4)
    label(legend_x+94, legend_y+84, "OIB", 23, "lm")
    label(WIDTH/2, 937, "PM: McDonough & Sun (1989); N-MORB: Gale et al. (2013); OIB: Sun & McDonough (1989)", 20, color="#53605a")

    OUT_PNG.parent.mkdir(parents=True, exist_ok=True)
    image.save(OUT_PNG, dpi=(300,300), optimize=True)
    document = canvas.Canvas(str(OUT_PDF), pagesize=(20*cm, 8.2*cm))
    document.drawImage(str(OUT_PNG), 0, 0, width=20*cm, height=8.2*cm, mask="auto")
    document.showPage()
    document.save()
    print(OUT_PNG)
    print(OUT_PDF)


if __name__ == "__main__":
    main()
